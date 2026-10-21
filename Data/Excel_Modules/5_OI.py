#!/usr/bin/env python3
"""
Data/Excel_Modules/5_OI.py
==========================
Dedicated Excel Module 5: Futures Open Interest (OI) Analysis.
Extracts real-time Futures Open Interest from CoinGlass (https://www.coinglass.com/open-interest/{asset})
across all 18 institutional universe assets:
BTC, ETH, XRP, SOL, BNB, DOGE, ADA, TRX, LINK, AVAX, SUI, NEAR, DOT, LTC, BCH, APT, OP, ARB.

Dual Institutional Layouts in Worksheet '5_OI':
  1. Open Interest Cross-Exchange Matrix (Assets in Rows, Exchanges in Columns):
     - Rows: All 18 institutional assets.
     - Columns: Asset, Total OI (USD), Binance, Bybit, OKX, CME, KuCoin, Gate,
       Bitget, BingX, MEXC, Bitunix, Hyperliquid, Others, Total OI (Tokens),
       24h Total Change (%), 4h Total Change (%), 1h Total Change (%), OI/24h Vol, Timestamp.
  2. Stacked Detailed Asset Table:
     - All assets stacked vertically with complete exchange-level granularity:
       Asset, Rank, Exchange, OI (Tokens), OI (USD), Market Share (%),
       OI Change (1h), OI Change (4h), OI Change (24h), OI/24h Vol, Timestamp.

Tab Hygiene & Shield:
  - Ephemeral CDP tab on port 9222, closed in finally block (arena.ai preserved).
  - Cloudflare Anycast high-speed routing.
"""

from __future__ import annotations

import argparse
import asyncio
import datetime
import json
import logging
import os
import pathlib
import sys
import time
import urllib.request

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl is required. Run 'pip install openpyxl'.")

try:
    import websockets
except ImportError:
    websockets = None

try:
    from cloudflare_rotator import rotate_ip
except ImportError:
    rotate_ip = None

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "Data"
DEFAULT_EXCEL_PATH = DATA_DIR / "DataBase_Trading.xlsx"
SHEET_NAME = "5_OI"
CDP_BASE_URL = "http://127.0.0.1:9222"

# 18 Institutional Binance USDT-M Perpetuals Universe
DEFAULT_ASSETS = [
    "BTC", "ETH", "XRP", "SOL", "BNB", "DOGE", "ADA", "TRX", "LINK",
    "AVAX", "SUI", "NEAR", "DOT", "LTC", "BCH", "APT", "OP", "ARB"
]

# Tracked Major Exchanges for Matrix View
MAJOR_EXCHANGES = [
    "Binance", "Bybit", "OKX", "CME", "KuCoin", "Gate",
    "Bitget", "BingX", "MEXC", "Bitunix", "Hyperliquid"
]

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] [5_OI] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S"
)
log = logging.getLogger("5_OI")


def open_ephemeral_tab(initial_url: str = "https://www.coinglass.com/open-interest/BTC") -> tuple[str, str]:
    """Creates a new dedicated ephemeral tab via CDP."""
    req = urllib.request.Request(f"{CDP_BASE_URL}/json/new?{initial_url}", method="PUT")
    new_tab = json.loads(urllib.request.urlopen(req, timeout=8).read().decode())
    return new_tab["id"], new_tab["webSocketDebuggerUrl"]


def close_tab(tab_id: str):
    """Closes tab via CDP."""
    try:
        urllib.request.urlopen(f"{CDP_BASE_URL}/json/close/{tab_id}", timeout=4)
        log.info(f"Closed ephemeral tab {tab_id}")
    except Exception as e:
        log.warning(f"Error closing tab {tab_id}: {e}")


def parse_usd_value(val_str: str) -> float:
    """Parses financial shorthand strings ($2.84B, $500.04M, $12.3K) into float."""
    if not val_str or val_str == "-":
        return 0.0
    s = val_str.replace("$", "").replace(",", "").strip()
    try:
        if s.endswith("B"):
            return float(s[:-1]) * 1e9
        elif s.endswith("M"):
            return float(s[:-1]) * 1e6
        elif s.endswith("K"):
            return float(s[:-1]) * 1e3
        return float(s)
    except Exception:
        return 0.0


async def extract_asset_open_interest(ws, asset: str) -> list[dict]:
    """Navigates to CoinGlass asset OI page and extracts all exchange rows."""
    target_url = f"https://www.coinglass.com/open-interest/{asset}"
    
    # 1. Navigate to asset page
    await ws.send(json.dumps({
        "id": 10,
        "method": "Page.navigate",
        "params": {"url": target_url}
    }))
    await ws.recv()

    # 2. Dynamic readiness check: wait for real exchange rows (not skeleton/measure rows)
    readiness_js = """
    (() => {
        const table = document.querySelector('.ant-table-body table') || 
                      Array.from(document.querySelectorAll('table')).find(t => t.querySelectorAll('tbody tr').length > 0);
        if (!table) return 0;
        const trs = Array.from(table.querySelectorAll('tbody tr:not(.ant-table-measure-row)'));
        const valid = trs.filter(tr => {
            const tds = tr.querySelectorAll('td');
            return tds.length >= 7 && tds[1].innerText.trim().length > 0;
        });
        return valid.length;
    })()
    """
    valid_count = 0
    for _ in range(40):  # up to 10 seconds
        await ws.send(json.dumps({"id": 11, "method": "Runtime.evaluate", "params": {"expression": readiness_js, "returnByValue": True}}))
        res = json.loads(await ws.recv())
        valid_count = res.get("result", {}).get("result", {}).get("value", 0)
        if valid_count >= 5:
            break
        await asyncio.sleep(0.25)

    # Auto-recovery: if table is still in skeleton state after 10s (possible 429), rotate IP & reload
    if valid_count < 3 and rotate_ip:
        log.warning(f"[{asset}] Table stuck on skeleton after 10s. Triggering Cloudflare Anycast IP rotation...")
        rotate_ip()
        await ws.send(json.dumps({"id": 13, "method": "Page.reload"}))
        await ws.recv()
        for _ in range(40):  # wait up to 10s after reload
            await ws.send(json.dumps({"id": 14, "method": "Runtime.evaluate", "params": {"expression": readiness_js, "returnByValue": True}}))
            res = json.loads(await ws.recv())
            valid_count = res.get("result", {}).get("result", {}).get("value", 0)
            if valid_count >= 5:
                break
            await asyncio.sleep(0.25)

    # 3. Extract table rows
    extract_js = """
    (() => {
        const table = document.querySelector('.ant-table-body table') || 
                      Array.from(document.querySelectorAll('table')).find(t => t.querySelectorAll('tbody tr').length > 0);
        if (!table) return [];
        const trs = Array.from(table.querySelectorAll('tbody tr'));
        return trs.map(tr => {
            const tds = Array.from(tr.querySelectorAll('td'));
            if (tds.length < 8) return null;

            const rank = tds[0].innerText.trim();
            const exchange = tds[1].innerText.trim();
            if (!exchange) return null;

            const oiTokens = tds[2] ? tds[2].innerText.replace(/\\n/g, ' ').trim() : "";
            const oiUsd = tds[3] ? tds[3].innerText.replace(/\\n/g, ' ').trim() : "";
            const rate = tds[4] ? tds[4].innerText.replace(/\\n/g, ' ').trim() : "";
            const chg1h = tds[5] ? tds[5].innerText.replace(/\\n/g, ' ').trim() : "";
            const chg4h = tds[6] ? tds[6].innerText.replace(/\\n/g, ' ').trim() : "";
            const chg24h = tds[7] ? tds[7].innerText.replace(/\\n/g, ' ').trim() : "";
            const oiVol = tds[8] ? tds[8].innerText.replace(/\\n/g, ' ').trim() : "";

            return {
                rank: rank || "All",
                exchange: exchange,
                oi_tokens: oiTokens,
                oi_usd: oiUsd,
                rate: rate,
                chg_1h: chg1h,
                chg_4h: chg4h,
                chg_24h: chg24h,
                oi_vol: oiVol
            };
        }).filter(r => r !== null);
    })()
    """
    await ws.send(json.dumps({"id": 12, "method": "Runtime.evaluate", "params": {"expression": extract_js, "returnByValue": True}}))
    res = json.loads(await ws.recv())
    return res.get("result", {}).get("result", {}).get("value", [])


async def harvest_universe_open_interest(assets: list[str] = DEFAULT_ASSETS) -> tuple[list[dict], list[dict]]:
    """
    Iterates through all 18 institutional assets in a single ephemeral tab.
    Returns:
      (matrix_records, stacked_records)
    """
    if not websockets:
        raise RuntimeError("websockets package is required.")

    tab_id, ws_url = open_ephemeral_tab(f"https://www.coinglass.com/open-interest/{assets[0]}")
    log.info(f"Opened dedicated ephemeral tab (ID: {tab_id}) for Open Interest harvest across {len(assets)} assets.")

    now_dt = datetime.datetime.now(datetime.timezone.utc)
    ts_str = now_dt.strftime("%Y-%m-%d %H:%M:%S UTC")

    matrix_rows: list[dict] = []
    stacked_rows: list[dict] = []

    try:
        ws = None
        for attempt in range(10):
            try:
                ws = await websockets.connect(ws_url, max_size=25 * 1024 * 1024, open_timeout=10)
                break
            except Exception:
                await asyncio.sleep(0.5)

        if not ws:
            raise RuntimeError(f"Could not connect to WebSocket at {ws_url}")

        async with ws:
            await asyncio.sleep(2.0)

            for i, asset in enumerate(assets, start=1):
                t0 = time.perf_counter()
                raw_rows = await extract_asset_open_interest(ws, asset)
                elapsed = time.perf_counter() - t0

                log.info(f"[{i:>2}/{len(assets)}] {asset:<5}: Extracted {len(raw_rows)} exchange rows in {elapsed:.2f}s")

                # Parse rows for Stacked Table & Matrix
                asset_exchanges: dict[str, str] = {}
                all_summary = None

                for r in raw_rows:
                    r["asset"] = asset
                    r["timestamp_utc"] = ts_str
                    stacked_rows.append(r)

                    exch = r.get("exchange", "")
                    if exch == "All" or r.get("rank") == "All":
                        all_summary = r
                    else:
                        asset_exchanges[exch] = r.get("oi_usd", "-")

                # Build Matrix Record for this Asset
                total_usd = all_summary.get("oi_usd", "-") if all_summary else "-"
                total_tokens = all_summary.get("oi_tokens", "-") if all_summary else "-"
                chg_24h = all_summary.get("chg_24h", "-") if all_summary else "-"
                chg_4h = all_summary.get("chg_4h", "-") if all_summary else "-"
                chg_1h = all_summary.get("chg_1h", "-") if all_summary else "-"
                oi_vol = all_summary.get("oi_vol", "-") if all_summary else "-"

                # Sum Others (exchanges not in MAJOR_EXCHANGES)
                major_sum = sum(parse_usd_value(asset_exchanges.get(e, "0")) for e in MAJOR_EXCHANGES)
                total_val = parse_usd_value(total_usd)
                others_val = max(0.0, total_val - major_sum)
                others_str = f"${others_val / 1e6:.2f}M" if others_val >= 1e6 else f"${others_val / 1e3:.2f}K" if others_val > 0 else "-"

                matrix_row = {
                    "asset": asset,
                    "total_oi_usd": total_usd,
                    "total_oi_tokens": total_tokens,
                    "chg_24h": chg_24h,
                    "chg_4h": chg_4h,
                    "chg_1h": chg_1h,
                    "oi_vol": oi_vol,
                    "others": others_str,
                    "timestamp_utc": ts_str
                }
                for exch in MAJOR_EXCHANGES:
                    matrix_row[exch] = asset_exchanges.get(exch, "-")

                matrix_rows.append(matrix_row)

    finally:
        close_tab(tab_id)

    log.info(
        f"Harvest complete: {len(matrix_rows)} matrix records and {len(stacked_rows)} detailed stacked records collected."
    )
    return matrix_rows, stacked_rows


def dump_to_excel(
    matrix_rows: list[dict],
    stacked_rows: list[dict],
    excel_path: pathlib.Path = DEFAULT_EXCEL_PATH
) -> bool:
    """Dumps both Matrix and Stacked sections into worksheet '5_OI' with institutional styling."""
    if not excel_path.exists():
        log.error(f"Excel workbook not found at: {excel_path}")
        return False

    if len(stacked_rows) < 50:
        log.error(
            f"Integrity check failed: only {len(stacked_rows)} exchange rows harvested across the universe. "
            f"Aborting Excel dump to prevent overwriting sheet '{SHEET_NAME}' with incomplete/empty data."
        )
        return False

    wb = openpyxl.load_workbook(excel_path)

    # Place sheet 5_OI at Index 6 (after 4_Hyperliquid_Analysis)
    if SHEET_NAME in wb.sheetnames:
        idx = wb.sheetnames.index(SHEET_NAME)
        wb.remove(wb[SHEET_NAME])
        ws = wb.create_sheet(title=SHEET_NAME, index=idx)
    else:
        target_idx = 6 if len(wb.sheetnames) >= 6 else len(wb.sheetnames)
        ws = wb.create_sheet(title=SHEET_NAME, index=target_idx)

    log.info(f"Target worksheet '{SHEET_NAME}' positioned at index {wb.sheetnames.index(SHEET_NAME)}")

    # Visual Styles
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    matrix_hdr_fill = PatternFill(start_color="161B22", end_color="161B22", fill_type="solid")
    stacked_hdr_fill = PatternFill(start_color="21262D", end_color="21262D", fill_type="solid")
    banner_fill = PatternFill(start_color="0D1117", end_color="0D1117", fill_type="solid")
    all_row_fill = PatternFill(start_color="F0F6FC", end_color="F0F6FC", fill_type="solid")

    base_font = Font(name="Segoe UI", size=9, color="000000")
    bold_font = Font(name="Segoe UI", size=9, bold=True, color="000000")
    green_font = Font(name="Segoe UI", size=9, bold=True, color="00A86B")
    red_font = Font(name="Segoe UI", size=9, bold=True, color="E53935")

    thin_border_side = Side(border_style="thin", color="E0E0E0")
    row_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # =========================================================================
    # SECTION 1: OPEN INTEREST MATRIX (Assets in Rows, Exchanges in Columns)
    # =========================================================================
    matrix_cols = [
        ("asset", "Asset", 12, "center"),
        ("total_oi_usd", "Total OI (USD)", 18, "right"),
    ] + [
        (exch, exch, 16, "right") for exch in MAJOR_EXCHANGES
    ] + [
        ("others", "Others", 14, "right"),
        ("total_oi_tokens", "Total OI (Tokens)", 20, "right"),
        ("chg_24h", "24h Total Change", 16, "right"),
        ("chg_4h", "4h Total Change", 16, "right"),
        ("chg_1h", "1h Total Change", 16, "right"),
        ("oi_vol", "OI / 24h Vol", 14, "right"),
        ("timestamp_utc", "Timestamp (UTC)", 24, "center"),
    ]

    # Section 1 Title Banner
    ws.merge_cells("A1:U1")
    banner_cell = ws.cell(row=1, column=1)
    banner_cell.value = "FUTURES OPEN INTEREST MATRIX - ASSETS (ROWS) x EXCHANGES (COLUMNS)"
    banner_cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    banner_cell.fill = banner_fill
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Section 1 Headers
    m_headers = [col[1] for col in matrix_cols]
    ws.append(m_headers)
    for c_idx in range(1, len(m_headers) + 1):
        cell = ws.cell(row=2, column=c_idx)
        cell.font = header_font
        cell.fill = matrix_hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 26

    # Section 1 Data Rows
    curr_row = 3
    for m_rec in matrix_rows:
        row_vals = [m_rec.get(col[0], "-") for col in matrix_cols]
        ws.append(row_vals)

        for c_idx, col_def in enumerate(matrix_cols, start=1):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.border = row_border
            key = col_def[0]
            val = str(cell.value)

            if key in ["chg_24h", "chg_4h", "chg_1h"]:
                cell.font = green_font if "+" in val else red_font if "-" in val else base_font
            elif key in ["asset", "total_oi_usd"]:
                cell.font = bold_font
            else:
                cell.font = base_font

            cell.alignment = Alignment(horizontal=col_def[3], vertical="center")

        ws.row_dimensions[curr_row].height = 20
        curr_row += 1

    # Blank Separator
    curr_row += 1

    # =========================================================================
    # SECTION 2: COMPREHENSIVE STACKED GRANULAR DATASET
    # =========================================================================
    stacked_cols = [
        ("asset", "Asset", 12, "center"),
        ("rank", "Rank", 8, "center"),
        ("exchange", "Exchange", 18, "left"),
        ("oi_tokens", "OI (Tokens)", 22, "right"),
        ("oi_usd", "OI (USD)", 18, "right"),
        ("rate", "Market Share (%)", 16, "right"),
        ("chg_1h", "OI Change (1h)", 16, "right"),
        ("chg_4h", "OI Change (4h)", 16, "right"),
        ("chg_24h", "OI Change (24h)", 16, "right"),
        ("oi_vol", "OI / 24h Vol", 14, "right"),
        ("timestamp_utc", "Timestamp (UTC)", 24, "center"),
    ]

    # Section 2 Title Banner
    banner2_row = curr_row
    ws.merge_cells(start_row=banner2_row, start_column=1, end_row=banner2_row, end_column=len(stacked_cols))
    banner2_cell = ws.cell(row=banner2_row, column=1)
    banner2_cell.value = "DETAILED OPEN INTEREST BREAKDOWN BY ASSET & EXCHANGE (STACKED GRANULAR DATASET)"
    banner2_cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    banner2_cell.fill = banner_fill
    banner2_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[banner2_row].height = 28
    curr_row += 1

    # Section 2 Headers
    s_headers = [col[1] for col in stacked_cols]
    ws.append(s_headers)
    for c_idx in range(1, len(s_headers) + 1):
        cell = ws.cell(row=curr_row, column=c_idx)
        cell.font = header_font
        cell.fill = stacked_hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[curr_row].height = 26
    curr_row += 1

    # Section 2 Data Rows
    for s_rec in stacked_rows:
        row_vals = [s_rec.get(col[0], "-") for col in stacked_cols]
        ws.append(row_vals)

        is_all_row = s_rec.get("exchange") == "All" or s_rec.get("rank") == "All"

        for c_idx, col_def in enumerate(stacked_cols, start=1):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.border = row_border
            key = col_def[0]
            val = str(cell.value)

            if is_all_row:
                cell.fill = all_row_fill

            if key in ["chg_1h", "chg_4h", "chg_24h"]:
                cell.font = green_font if "+" in val else red_font if "-" in val else base_font
            elif key in ["asset", "exchange"] or is_all_row:
                cell.font = bold_font
            else:
                cell.font = base_font

            cell.alignment = Alignment(horizontal=col_def[3], vertical="center")

        ws.row_dimensions[curr_row].height = 20
        curr_row += 1

    # Apply Column Widths
    for c_idx, col_def in enumerate(matrix_cols, start=1):
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = col_def[2]

    ws.views.sheetView[0].showGridLines = True
    wb.save(excel_path)
    log.info(
        f"Successfully dumped Open Interest data ({len(matrix_rows)} matrix assets, {len(stacked_rows)} stacked rows) into '{SHEET_NAME}' at {excel_path}."
    )
    return True


def main():
    parser = argparse.ArgumentParser(description="CoinGlass Futures Open Interest Matrix & Stacked Dump to Excel")
    parser.add_argument("--assets", nargs="+", default=DEFAULT_ASSETS, help="List of assets to harvest")
    parser.add_argument("--excel-path", type=pathlib.Path, default=DEFAULT_EXCEL_PATH, help="Path to DataBase_Trading.xlsx")
    args = parser.parse_args()

    matrix_rows, stacked_rows = asyncio.run(harvest_universe_open_interest(assets=args.assets))
    if not matrix_rows and not stacked_rows:
        log.error("No Open Interest records harvested.")
        sys.exit(1)

    ok = dump_to_excel(matrix_rows, stacked_rows, excel_path=args.excel_path)
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
