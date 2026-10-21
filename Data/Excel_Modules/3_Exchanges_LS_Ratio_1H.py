#!/usr/bin/env python3
"""
Data/Excel_Modules/3_Exchanges_LS_Ratio_1H.py
============================================
Dedicated Excel Module 3B: Exchanges Long/Short Ratio (1 Hour).
Extracts live Exchanges Long/Short Ratio and Taker Buy/Sell Volume across
the 18 institutional perpetual assets in our trading universe:
  BTC, ETH, XRP, SOL, BNB, DOGE, ADA, TRX, LINK, AVAX, SUI, NEAR, DOT, LTC, BCH, APT, OP, ARB

Target URL: https://www.coinglass.com/LongShortRatio/{asset}
Timeframe: 1 hour (toggled in UI)
Worksheet: '3_Exchanges_LS_Ratio_1H' in Data/DataBase_Trading.xlsx
Position: Immediately following '3_Exchanges_LS_Ratio_15M'

Browser Hygiene:
  - Opens dedicated ephemeral tab over CDP port 9222.
  - Navigates through universe assets.
  - Always closes the tab in a finally block upon completion.
"""

from __future__ import annotations

import argparse
import asyncio
import datetime
import json
import logging
import os
import pathlib
import sys
import time
import urllib.request

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl is required. Run 'pip install openpyxl'.")

try:
    import websockets
except ImportError:
    websockets = None

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "Data"
DEFAULT_EXCEL_PATH = DATA_DIR / "DataBase_Trading.xlsx"
CACHE_JSON_PATH = DATA_DIR / "coinglass" / "latest" / "exchanges_ls_ratio_1h.json"
CDP_BASE_URL = "http://127.0.0.1:9222"

SHEET_NAME = "3_Exchanges_LS_Ratio_1H"
REFERENCE_SHEET = "3_Exchanges_LS_Ratio_15M"
TIMEFRAME_LABEL = "1 hour"
TIMEFRAME_CODE = "1h"

# Canonical 18 Institutional Binance USDT-M Perpetuals Universe
DEFAULT_UNIVERSE = [
    "BTC", "ETH", "XRP", "SOL", "BNB", "DOGE", "ADA", "TRX", "LINK", 
    "AVAX", "SUI", "NEAR", "DOT", "LTC", "BCH", "APT", "OP", "ARB"
]

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] [3_LS_Ratio_1H] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S"
)
log = logging.getLogger("3_LS_Ratio_1H")

COLUMNS = [
    ("Asset", "Asset", 12, "center"),
    ("Rank", "#", 8, "center"),
    ("Exchange", "Exchange", 18, "left"),
    ("Long_Pct", "Long %", 14, "right"),
    ("Short_Pct", "Short %", 14, "right"),
    ("LS_Ratio", "L/S Ratio", 14, "right"),
    ("Long_Vol", "Long Volume", 18, "right"),
    ("Short_Vol", "Short Volume", 18, "right"),
    ("Total_Vol", "Total Volume", 18, "right"),
    ("Timeframe", "Timeframe", 14, "center"),
    ("Timestamp_UTC", "Timestamp (UTC)", 24, "center")
]


def open_ephemeral_tab(initial_url: str = "https://www.coinglass.com/LongShortRatio/BTC") -> tuple[str, str]:
    """Creates a new dedicated ephemeral tab via CDP."""
    req = urllib.request.Request(f"{CDP_BASE_URL}/json/new?{initial_url}", method="PUT")
    new_tab = json.loads(urllib.request.urlopen(req, timeout=5).read().decode())
    return new_tab["id"], new_tab["webSocketDebuggerUrl"]


def close_tab(tab_id: str):
    """Closes tab via CDP."""
    try:
        urllib.request.urlopen(f"{CDP_BASE_URL}/json/close/{tab_id}", timeout=4)
        log.info(f"Closed tab {tab_id}")
    except Exception as e:
        log.warning(f"Error closing tab: {e}")


def calculate_ls_ratio(long_pct_str: str, short_pct_str: str) -> str:
    """Calculates numerical Long/Short Ratio from percentage strings."""
    try:
        lp = float(long_pct_str.replace("%", "").strip())
        sp = float(short_pct_str.replace("%", "").strip())
        if sp > 0:
            return f"{lp / sp:.2f}"
    except Exception:
        pass
    return "-"


def parse_volume_num(vol_str: str) -> float:
    """Parses volume string like $2.75B, $832.34M, $298.17K to numeric float."""
    try:
        s = vol_str.replace("$", "").replace(",", "").strip().upper()
        if not s or s == "-":
            return 0.0
        if s.endswith("B"):
            return float(s[:-1]) * 1e9
        elif s.endswith("M"):
            return float(s[:-1]) * 1e6
        elif s.endswith("K"):
            return float(s[:-1]) * 1e3
        return float(s)
    except Exception:
        return 0.0


def format_volume(num: float) -> str:
    """Formats float volume into standard financial notation ($B, $M, $K)."""
    if num >= 1e9:
        return f"${num / 1e9:.2f}B"
    elif num >= 1e6:
        return f"${num / 1e6:.2f}M"
    elif num >= 1e3:
        return f"${num / 1e3:.2f}K"
    elif num > 0:
        return f"${num:.2f}"
    return "$0"


async def extract_asset_ls_ratio(ws, asset: str, target_timeframe: str = TIMEFRAME_LABEL) -> list[dict]:
    """Navigates to asset page, verifies/switches timeframe, and extracts all exchange rows."""
    target_url = f"https://www.coinglass.com/LongShortRatio/{asset}"
    
    # 1. Navigate to asset page
    await ws.send(json.dumps({
        "id": 10,
        "method": "Page.navigate",
        "params": {"url": target_url}
    }))
    await ws.recv()

    # 2. Wait for page DOM container to stabilize
    container_ready_js = """
    (() => {
        const headings = Array.from(document.querySelectorAll('*')).filter(el => {
            return el.children.length === 0 && el.innerText && el.innerText.trim().startsWith('Exchanges') && el.innerText.includes('Long/Short Ratio');
        });
        if (headings.length === 0) return false;
        let container = headings[0].parentElement;
        while (container && (!container.innerText.includes('Binance') || !container.querySelector('button.MuiSelect-button'))) {
            container = container.parentElement;
        }
        if (!container) return false;
        const rows = Array.from(container.querySelectorAll('.ant-row'));
        if (!rows.length) return false;
        return rows[0].children.length >= 15;
    })()
    """
    is_ready = False
    for _ in range(30):
        await asyncio.sleep(0.1)
        await ws.send(json.dumps({"id": 11, "method": "Runtime.evaluate", "params": {"expression": container_ready_js, "returnByValue": True}}))
        res = json.loads(await ws.recv())
        if res.get("result", {}).get("result", {}).get("value", False):
            is_ready = True
            break

    if not is_ready:
        await asyncio.sleep(1.0)

    # 3. Timeframe Verification & Toggle
    check_and_toggle_tf_js = f"""
    (() => {{
        const headings = Array.from(document.querySelectorAll('*')).filter(el => {{
            return el.children.length === 0 && el.innerText && el.innerText.trim().startsWith('Exchanges') && el.innerText.includes('Long/Short Ratio');
        }});
        if (headings.length === 0) return {{status: 'no_heading'}};
        let container = headings[0].parentElement;
        while (container && (!container.innerText.includes('Binance') || !container.querySelector('button.MuiSelect-button'))) {{
            container = container.parentElement;
        }}
        if (!container) return {{status: 'no_container'}};

        const tfBtn = container.querySelector('button.MuiSelect-button');
        if (!tfBtn) return {{status: 'no_button'}};

        const currentTf = tfBtn.innerText.trim().toLowerCase();
        if (currentTf === '{target_timeframe.lower()}') {{
            return {{status: 'already_match', current: currentTf}};
        }}

        const listboxId = tfBtn.getAttribute('aria-controls');
        tfBtn.click();
        return {{status: 'clicked_open', listboxId: listboxId}};
    }})()
    """
    await ws.send(json.dumps({"id": 20, "method": "Runtime.evaluate", "params": {"expression": check_and_toggle_tf_js, "returnByValue": True}}))
    tf_res = json.loads(await ws.recv())
    tf_val = tf_res.get("result", {}).get("result", {}).get("value", {})

    if tf_val.get("status") == "clicked_open":
        await asyncio.sleep(0.3)
        listbox_id = tf_val.get("listboxId")
        click_option_js = f"""
        (() => {{
            const listbox = document.getElementById('{listbox_id}') || document.querySelector('.MuiMenu-list');
            if (!listbox) return false;
            const opts = Array.from(listbox.querySelectorAll('[role="option"], li'));
            const targetOpt = opts.find(o => o.innerText.trim().toLowerCase() === '{target_timeframe.lower()}');
            if (targetOpt) {{
                targetOpt.click();
                return true;
            }}
            return false;
        }})()
        """
        await ws.send(json.dumps({"id": 21, "method": "Runtime.evaluate", "params": {"expression": click_option_js, "returnByValue": True}}))
        await ws.recv()
        await asyncio.sleep(0.8)

    # 4. Scoped Exchange Row Extraction
    extract_rows_js = """
    (() => {
        const headings = Array.from(document.querySelectorAll('*')).filter(el => {
            return el.children.length === 0 && el.innerText && el.innerText.trim().startsWith('Exchanges') && el.innerText.includes('Long/Short Ratio');
        });
        if (headings.length === 0) return [];
        let container = headings[0].parentElement;
        while (container && (!container.innerText.includes('Binance') || !container.querySelector('button.MuiSelect-button'))) {
            container = container.parentElement;
        }
        if (!container) return [];

        const rows = Array.from(container.querySelectorAll('.ant-row')).filter(r => {
            const txt = r.innerText || '';
            return (txt.includes('Long') && txt.includes('Short')) || (txt.includes('%') && txt.includes('$'));
        });
        if (!rows.length) return [];

        const cols = Array.from(rows[0].children).map(c => c.innerText.trim());
        const records = [];

        for (let i = 0; i < cols.length; i += 3) {
            if (i + 2 >= cols.length) break;

            const exchangeBlock = cols[i].split('\\n');
            const rank = exchangeBlock.length > 1 ? exchangeBlock[0].trim() : '-';
            const exchange = exchangeBlock.length > 1 ? exchangeBlock[1].trim() : exchangeBlock[0].trim();

            const pctBlock = cols[i+1].split('\\n');
            const longPct = pctBlock[0] ? pctBlock[0].trim() : '-';
            const shortPct = pctBlock[1] ? pctBlock[1].trim() : '-';

            const volBlock = cols[i+2].split('\\n');
            let longVol = '-';
            let shortVol = '-';
            for (let j = 0; j < volBlock.length; j++) {
                if (volBlock[j].trim() === 'Long' && j + 1 < volBlock.length) longVol = volBlock[j+1].trim();
                if (volBlock[j].trim() === 'Short' && j + 1 < volBlock.length) shortVol = volBlock[j+1].trim();
            }

            records.push({
                rank: rank,
                exchange: exchange,
                long_pct: longPct,
                short_pct: shortPct,
                long_vol: longVol,
                short_vol: shortVol
            });
        }
        return records;
    })()
    """
    await ws.send(json.dumps({"id": 30, "method": "Runtime.evaluate", "params": {"expression": extract_rows_js, "returnByValue": True}}))
    rows_res = json.loads(await ws.recv())
    raw_records = rows_res.get("result", {}).get("result", {}).get("value", [])

    now_dt = datetime.datetime.now(datetime.timezone.utc)
    ts_str = now_dt.strftime("%Y-%m-%d %H:%M:%S UTC")

    parsed_records = []
    for r in raw_records:
        lp = r.get("long_pct", "-")
        sp = r.get("short_pct", "-")
        lv = r.get("long_vol", "-")
        sv = r.get("short_vol", "-")
        ls_ratio = calculate_ls_ratio(lp, sp)

        lv_num = parse_volume_num(lv)
        sv_num = parse_volume_num(sv)
        tv_str = format_volume(lv_num + sv_num) if (lv_num > 0 or sv_num > 0) else "-"

        rank_val = "Total" if r.get("rank") == "-" else r.get("rank", "-")

        parsed_records.append({
            "Asset": asset,
            "Rank": rank_val,
            "Exchange": r.get("exchange", "-"),
            "Long_Pct": lp,
            "Short_Pct": sp,
            "LS_Ratio": ls_ratio,
            "Long_Vol": lv,
            "Short_Vol": sv,
            "Total_Vol": tv_str,
            "Timeframe": TIMEFRAME_CODE,
            "Timestamp_UTC": ts_str
        })

    return parsed_records


async def harvest_universe_ls_ratio(universe: list[str] = DEFAULT_UNIVERSE, target_timeframe: str = TIMEFRAME_LABEL) -> list[dict]:
    """Iterates across universe assets sequentially inside a single ephemeral tab."""
    if not websockets:
        log.error("websockets package required for CDP connection.")
        return []

    initial_url = f"https://www.coinglass.com/LongShortRatio/{universe[0]}"
    tab_id, ws_url = open_ephemeral_tab(initial_url)
    log.info(f"Opened dedicated ephemeral tab (ID: {tab_id}) for {TIMEFRAME_CODE} harvest across {len(universe)} assets.")

    all_harvested = []
    start_t = time.perf_counter()

    try:
        await asyncio.sleep(3.0)
        async with websockets.connect(ws_url, max_size=25 * 1024 * 1024, open_timeout=10) as ws:
            for idx, asset in enumerate(universe, 1):
                asset_start = time.perf_counter()
                records = await extract_asset_ls_ratio(ws, asset=asset, target_timeframe=target_timeframe)
                elapsed_asset = time.perf_counter() - asset_start

                if records:
                    all_harvested.extend(records)
                    log.info(f"[{idx:>2}/{len(universe)}] {asset:<5}: Extracted {len(records)} exchange rows in {elapsed_asset:.2f}s")
                else:
                    log.warning(f"[{idx:>2}/{len(universe)}] {asset:<5}: No records extracted (took {elapsed_asset:.2f}s)")

        total_elapsed = time.perf_counter() - start_t
        log.info(f"Harvest complete: {len(all_harvested)} total records across {len(universe)} assets in {total_elapsed:.2f}s.")

        CACHE_JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
        with open(CACHE_JSON_PATH, "w", encoding="utf-8") as f:
            json.dump({
                "timeframe": TIMEFRAME_CODE,
                "timeframe_label": target_timeframe,
                "harvest_timestamp_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
                "total_records": len(all_harvested),
                "universe_count": len(universe),
                "records": all_harvested
            }, f, indent=2)

        return all_harvested
    except Exception as e:
        log.error(f"Error during {TIMEFRAME_CODE} Long/Short Ratio harvest: {e}")
        return all_harvested
    finally:
        close_tab(tab_id)


def dump_ls_ratio_to_excel(records: list[dict], excel_path: pathlib.Path = DEFAULT_EXCEL_PATH) -> bool:
    """Writes harvested Long/Short Ratio records to Excel worksheet with full heatmaps."""
    if not records:
        log.error(f"No records available to dump for {SHEET_NAME}.")
        return False

    if not excel_path.exists():
        log.error(f"Target Excel file does not exist: {excel_path}")
        return False

    wb = openpyxl.load_workbook(excel_path)

    # Position sheet following REFERENCE_SHEET
    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        if REFERENCE_SHEET in wb.sheetnames:
            ref_idx = wb.sheetnames.index(REFERENCE_SHEET)
            ws = wb.create_sheet(title=SHEET_NAME, index=ref_idx + 1)
        elif "2_Spot_Netflow_S" in wb.sheetnames:
            ref_idx = wb.sheetnames.index("2_Spot_Netflow_S")
            ws = wb.create_sheet(title=SHEET_NAME, index=ref_idx + 1)
        else:
            ws = wb.create_sheet(title=SHEET_NAME, index=3)

    log.info(f"Target worksheet '{SHEET_NAME}' positioned at index {wb.sheetnames.index(SHEET_NAME)}")

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    asset_total_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    asset_total_font = Font(name="Segoe UI", size=10, bold=True, color="0369A1")

    exchange_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    data_font = Font(name="Segoe UI", size=10, bold=False, color="1E293B")
    dim_font = Font(name="Segoe UI", size=10, bold=False, color="64748B")

    green_font = Font(name="Segoe UI", size=10, bold=True, color="15803D")
    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    red_font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    neutral_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # 1. Write Header Row
    for col_idx, (key, display_name, width, align) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=col_idx, value=display_name)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal=align, vertical="center")
        c.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(width, 12)
    ws.row_dimensions[1].height = 26

    # 2. Freeze Header Pane
    ws.freeze_panes = "A2"

    # 3. Write Data Rows
    for idx, rec in enumerate(records):
        r_idx = 2 + idx
        is_total_row = (rec.get("Rank") == "Total")

        lp_str = rec.get("Long_Pct", "")
        sp_str = rec.get("Short_Pct", "")
        lp_val = float(lp_str.replace("%", "").strip()) if "%" in lp_str else 50.0
        sp_val = float(sp_str.replace("%", "").strip()) if "%" in sp_str else 50.0

        for col_idx, (key, _, _, align) in enumerate(COLUMNS, 1):
            val = rec.get(key, "-")
            cell_obj = ws.cell(row=r_idx, column=col_idx, value=val)
            cell_obj.border = thin_border
            cell_obj.alignment = Alignment(horizontal=align, vertical="center")

            if is_total_row:
                cell_obj.fill = asset_total_fill
                cell_obj.font = asset_total_font
            else:
                cell_obj.fill = neutral_fill
                if key in ["Asset", "Exchange"]:
                    cell_obj.font = exchange_font
                elif key == "Long_Pct":
                    if lp_val > sp_val:
                        cell_obj.fill = green_fill
                        cell_obj.font = green_font
                    else:
                        cell_obj.font = dim_font
                elif key == "Short_Pct":
                    if sp_val > lp_val:
                        cell_obj.fill = red_fill
                        cell_obj.font = red_font
                    else:
                        cell_obj.font = dim_font
                elif key in ["Long_Vol", "Short_Vol", "Total_Vol"]:
                    cell_obj.font = data_font
                elif key == "LS_Ratio":
                    if lp_val > sp_val:
                        cell_obj.font = green_font
                    elif sp_val > lp_val:
                        cell_obj.font = red_font
                    else:
                        cell_obj.font = data_font
                else:
                    cell_obj.font = data_font

        ws.row_dimensions[r_idx].height = 20

    # 4. Prune trailing rows
    current_data_end = 1 + len(records)
    if ws.max_row > current_data_end:
        ws.delete_rows(current_data_end + 1, ws.max_row - current_data_end)

    # 5. Save Atomically
    temp_excel = excel_path.with_suffix(".tmp.xlsx")
    wb.save(temp_excel)
    wb.close()
    temp_excel.replace(excel_path)

    log.info(f"Successfully dumped {len(records)} rows into '{SHEET_NAME}' at {excel_path}.")
    return True


def run_ls_ratio_1h_module(universe: list[str] = DEFAULT_UNIVERSE, excel_path: pathlib.Path = DEFAULT_EXCEL_PATH) -> bool:
    """Main orchestration for Module 3B: Exchanges Long/Short Ratio (1H)."""
    records = asyncio.run(harvest_universe_ls_ratio(universe=universe, target_timeframe=TIMEFRAME_LABEL))
    if not records:
        log.error("Failed to acquire Long/Short Ratio records.")
        return False
    return dump_ls_ratio_to_excel(records, excel_path=excel_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Exchanges Long/Short Ratio 1H Excel Dumper")
    parser.add_argument("--symbols", type=str, default="", help="Comma-separated symbols override (e.g. BTC,ETH)")
    args = parser.parse_args()

    symbols_list = [s.strip().upper() for s in args.symbols.split(",") if s.strip()] if args.symbols else DEFAULT_UNIVERSE
    success = run_ls_ratio_1h_module(universe=symbols_list)
    sys.exit(0 if success else 1)
