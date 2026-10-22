#!/usr/bin/env python3
"""
Data/Excel_Modules/15_Hyperdash_Stops_Map.py
============================================
Hyperdash Multi-Asset On-Chain Stop-Loss Orders Price Ladder & Cumulative Map (100% Native API).
Captures point-by-point buy/sell stop-loss clusters and cumulative curves across all 18 institutional
perpetual assets directly from Hyperdash GraphQL API (https://api.hyperdash.com/graphql) without browser or CDP overhead,
and stacks them asset-over-asset into worksheet '15_Hyperdash_Stops_Map' in Data/DataBase_Trading.xlsx
(positioned at Index 23, immediately after '14_Hyperdash_Liq_Map').

Universe (18 Institutional Binance USDT-M Perpetuals):
  BTC, ETH, XRP, SOL, BNB, DOGE, ADA, TRX, LINK, AVAX, SUI, NEAR, DOT, LTC, BCH, APT, OP, ARB

Layout Matches Operator Specifications & 8_Exchange_Liquidation_Map_1D:
  Sr.No | Asset | Time | TimeStamp | Mark Price | Price | Distance (%) | Side | Cumulative Stops | Buy Stops | Sell Stops | Size (Coins)
"""

from __future__ import annotations

import argparse
import datetime
import json
import logging
import pathlib
import sys
import time
import urllib.request

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] [%(name)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger("15_Hyperdash_Stops_Map")

MODULES_DIR = pathlib.Path(__file__).resolve().parent
ROOT_DIR = MODULES_DIR.parents[1]
DEFAULT_EXCEL_PATH = ROOT_DIR / "Data" / "DataBase_Trading.xlsx"
GRAPHQL_URL = "https://api.hyperdash.com/graphql"
SHEET_NAME = "15_Hyperdash_Stops_Map"

DEFAULT_UNIVERSE = [
    "BTC", "ETH", "XRP", "SOL", "BNB",
    "DOGE", "ADA", "TRX", "LINK", "AVAX",
    "SUI", "NEAR", "DOT", "LTC", "BCH",
    "APT", "OP", "ARB",
]

COLUMNS_DEF = [
    ("sr_no", "Sr.No", 10, "center"),
    ("asset", "Asset", 12, "center"),
    ("time_window", "Time", 10, "center"),
    ("timestamp_utc", "TimeStamp", 24, "center"),
    ("mark_price", "Mark Price", 16, "right"),
    ("price", "Price", 16, "right"),
    ("distance_pct", "Distance (%)", 14, "right"),
    ("side", "Side", 14, "center"),
    ("cumulative_stops", "Cumulative Stops", 24, "right"),
    ("buy_stops", "Buy Stops", 20, "right"),
    ("sell_stops", "Sell Stops", 20, "right"),
    ("size_coins", "Size (Coins)", 18, "right"),
]

QUERY_STOPS = """
query GetStopOrderLevelsV2($coin: String!, $minPrice: Float!, $maxPrice: Float!, $startTime: Float!, $endTime: Float) {
  analytics {
    stopOrderLevels: stopOrderLevelsV2(
      coin: $coin
      minPrice: $minPrice
      maxPrice: $maxPrice
      startTime: $startTime
      endTime: $endTime
    ) {
      coin
      currentPrice
      bandSize
      bands {
        minPrice
        maxPrice
        historicalData {
          timestamp
          totalAmount
        }
      }
    }
  }
}
"""


def format_usd(val: float | None) -> str:
    """Formats float amount to human-readable string matching CoinGlass/Hyperdash UI."""
    if val is None or val == 0:
        return "$0"
    abs_val = abs(val)
    if abs_val >= 1e9:
        return f"${val / 1e9:.2f}B"
    elif abs_val >= 1e6:
        return f"${val / 1e6:.2f}M"
    elif abs_val >= 1e3:
        return f"${val / 1e3:.2f}K"
    elif abs_val >= 1:
        return f"${val:.2f}"
    else:
        return f"${val:.4f}"


def format_price(val: float | None) -> str:
    if val is None:
        return "$0"
    if val >= 1000:
        return f"${val:,.2f}"
    elif val >= 1:
        return f"${val:,.4f}"
    else:
        return f"${val:,.6f}"


def fetch_stops_ladder(coin: str, now_ms: float, start_ms: float) -> tuple[float, list[dict]]:
    """Fetches full stop order band ladder for coin via Hyperdash native GraphQL API."""
    payload = json.dumps({
        "query": QUERY_STOPS,
        "variables": {
            "coin": coin,
            "minPrice": 0.0,
            "maxPrice": 10000000.0,
            "startTime": start_ms,
            "endTime": now_ms
        }
    }).encode("utf-8")
    
    req = urllib.request.Request(
        GRAPHQL_URL,
        data=payload,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
            "Content-Type": "application/json"
        }
    )
    
    try:
        with urllib.request.urlopen(req, timeout=12) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            analytics = data.get("data", {}).get("analytics", {})
            stop_data = analytics.get("stopOrderLevels")
            if not stop_data:
                return 0.0, []
            
            mark_price = float(stop_data.get("currentPrice") or 0.0)
            bands = stop_data.get("bands", [])
            
            active_bands = []
            for b in bands:
                hist = b.get("historicalData")
                if hist and len(hist) > 0:
                    last_pt = hist[-1]
                    amt = float(last_pt.get("totalAmount") or 0.0)
                    if amt > 0:
                        mid_px = (float(b["minPrice"]) + float(b["maxPrice"])) / 2.0
                        notional = amt * mid_px
                        is_buy = mid_px > mark_price
                        active_bands.append({
                            "mid_px": mid_px,
                            "amt": amt,
                            "notional": notional,
                            "side": "Buy Stop" if is_buy else "Sell Stop",
                        })
            return mark_price, active_bands
    except Exception as e:
        log.warning(f"Error fetching stops ladder for {coin}: {e}")
        return 0.0, []


def harvest_all_stops_maps(symbols: list[str]) -> list[dict]:
    """Iterates through symbols, calculates cumulative curves outward from mark, and stacks rows."""
    all_records = []
    global_sr_no = 1
    now = datetime.datetime.now(datetime.timezone.utc)
    ts_str = now.strftime("%Y-%m-%d %H:%M:%S UTC")
    now_ms = now.timestamp() * 1000
    start_ms = now_ms - 24 * 3600 * 1000

    for idx, sym in enumerate(symbols, start=1):
        log.info(f"[{idx}/{len(symbols)}] Harvesting Hyperdash Stops Map for {sym}...")
        t0 = time.time()
        mark_px, bands = fetch_stops_ladder(sym, now_ms, start_ms)
        if not bands or mark_px <= 0:
            log.warning(f"  No active stop bands for {sym}")
            continue

        # Separate into Sell Stops (< mark_px) and Buy Stops (> mark_px)
        sells = [b for b in bands if b["side"] == "Sell Stop"]
        buys = [b for b in bands if b["side"] == "Buy Stop"]

        # Sort Sells descending by price (closest to mark first for cumulative sum)
        sells.sort(key=lambda x: x["mid_px"], reverse=True)
        cum_sells = 0.0
        for b in sells:
            cum_sells += b["notional"]
            b["cum_notional"] = cum_sells

        # Sort Buys ascending by price (closest to mark first for cumulative sum)
        buys.sort(key=lambda x: x["mid_px"])
        cum_buys = 0.0
        for b in buys:
            cum_buys += b["notional"]
            b["cum_notional"] = cum_buys

        # For the worksheet display: stack in price ascending order (institutional ladder)
        all_sym_bands = sells + buys
        all_sym_bands.sort(key=lambda x: x["mid_px"])

        log.info(f"  {sym}: Mark={mark_px:,.2f} | Sell Stops={len(sells)} | Buy Stops={len(buys)} | Total={len(all_sym_bands)} in {time.time()-t0:.2f}s")

        for b in all_sym_bands:
            px = b["mid_px"]
            side = b["side"]
            notional = b["notional"]
            cum_notional = b["cum_notional"]
            amt = b["amt"]
            dist_pct = ((px - mark_px) / mark_px) * 100.0

            buy_val = format_usd(notional) if side == "Buy Stop" else "$0"
            sell_val = format_usd(notional) if side == "Sell Stop" else "$0"
            cum_val = format_usd(cum_notional)

            all_records.append({
                "sr_no": global_sr_no,
                "asset": sym,
                "time_window": "Live",
                "timestamp_utc": ts_str,
                "mark_price": format_price(mark_px),
                "price": format_price(px),
                "distance_pct": f"{dist_pct:+.2f}%",
                "side": side,
                "cumulative_stops": cum_val,
                "buy_stops": buy_val,
                "sell_stops": sell_val,
                "size_coins": f"{amt:,.4f} {sym}",
            })
            global_sr_no += 1

    log.info(f"Total Stops Map records across all {len(symbols)} assets: {len(all_records)} rows.")
    return all_records


def dump_to_excel(records: list[dict], excel_path: pathlib.Path = DEFAULT_EXCEL_PATH) -> bool:
    """Dumps stacked stops map into worksheet '15_Hyperdash_Stops_Map' at Index 23."""
    if not excel_path.exists():
        log.error(f"Excel workbook not found at: {excel_path}")
        return False

    if not records:
        log.error("Integrity guard: 0 records extracted. Aborting dump to prevent empty sheet.")
        return False

    wb = openpyxl.load_workbook(excel_path)

    # Position at Index 23 (immediately after 14_Hyperdash_Liq_Map at Index 22)
    if SHEET_NAME in wb.sheetnames:
        idx = wb.sheetnames.index(SHEET_NAME)
        wb.remove(wb[SHEET_NAME])
        ws = wb.create_sheet(title=SHEET_NAME, index=idx)
    else:
        target_idx = 23 if len(wb.sheetnames) >= 23 else len(wb.sheetnames)
        ws = wb.create_sheet(title=SHEET_NAME, index=target_idx)

    log.info(f"Target worksheet '{SHEET_NAME}' positioned at index {wb.sheetnames.index(SHEET_NAME)}")

    # Visual Styles matching 8_Exchange_Liquidation_Map_1D
    hdr_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    base_font = Font(name="Segoe UI", size=9, color="000000")
    bold_font = Font(name="Segoe UI", size=9, bold=True, color="000000")
    cum_font = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
    green_font = Font(name="Segoe UI", size=9, bold=True, color="00A86B")
    red_font = Font(name="Segoe UI", size=9, bold=True, color="E53935")
    subdued_font = Font(name="Segoe UI", size=9, color="888888")

    thin_border_side = Side(border_style="thin", color="E2E8F0")
    row_border = Border(
        left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side
    )

    # 1. Header Row
    ws.row_dimensions[1].height = 26
    for col_idx, (col_key, header_title, width, align) in enumerate(COLUMNS_DEF, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_title)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.border = row_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 2. Data Rows
    for row_idx, r in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, (col_key, _, _, align) in enumerate(COLUMNS_DEF, start=1):
            val = r.get(col_key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(horizontal=align, vertical="center")
            cell.border = row_border

            # Typography and semantics
            if col_key in ("sr_no", "asset", "time_window"):
                cell.font = bold_font
            elif col_key == "cumulative_stops":
                cell.font = cum_font
            elif col_key == "buy_stops":
                cell.font = green_font if val != "$0" else subdued_font
            elif col_key == "sell_stops":
                cell.font = red_font if val != "$0" else subdued_font
            elif col_key == "side":
                cell.font = green_font if val == "Buy Stop" else red_font
            elif col_key == "distance_pct":
                cell.font = green_font if str(val).startswith("+") else (red_font if str(val).startswith("-") else base_font)
            else:
                cell.font = base_font

    ws.freeze_panes = "A2"

    temp_excel = excel_path.with_suffix(".tmp.xlsx")
    wb.save(temp_excel)
    wb.close()
    import zipfile
    with zipfile.ZipFile(temp_excel, "r") as _z:
        if _z.testzip() is None:
            temp_excel.replace(excel_path)
        else:
            if temp_excel.exists(): temp_excel.unlink()
            raise ValueError("Corrupt zip produced during export")
    log.info(f"Successfully saved {len(records)} rows to '{SHEET_NAME}' at {excel_path}")
    return True


def main():
    parser = argparse.ArgumentParser(description="Hyperdash 15_Hyperdash_Stops_Map Generator")
    parser.add_argument("--symbols", nargs="+", default=DEFAULT_UNIVERSE, help="Assets to harvest")
    parser.add_argument("--excel-path", type=pathlib.Path, default=DEFAULT_EXCEL_PATH, help="Path to Excel workbook")
    args = parser.parse_args()

    symbols = args.symbols
    if len(symbols) == 1 and "," in symbols[0]:
        symbols = [s.strip().upper() for s in symbols[0].split(",") if s.strip()]

    log.info(f"Starting Hyperdash Stops Map harvest for {len(symbols)} assets: {symbols}")
    records = harvest_all_stops_maps(symbols=symbols)

    if not records:
        log.error("Harvesting yielded 0 records. Excel dump aborted.")
        sys.exit(1)

    success = dump_to_excel(records, excel_path=args.excel_path)
    if not success:
        sys.exit(1)

    log.info("15_Hyperdash_Stops_Map pipeline executed successfully.")


if __name__ == "__main__":
    main()
