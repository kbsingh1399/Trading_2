#!/usr/bin/env python3
"""
Data/Excel_Modules/17_Orderbook.py
==================================
Hyperdash & Hyperliquid Multi-Asset L2 & L3 Orderbook Architecture & Whale Wallet Harvester.
Captures institutional orderbooks across all asset classes:
  1. CRYPTO: BTC, ETH, SOL, BNB, XRP, DOGE, AVAX, SUI
  2. INDICES: S&P 500 (xyz:SP500)
  3. METALS: Gold (xyz:GOLD), Silver (xyz:SILVER)
  4. COMMODITIES: Crude Oil WTI (xyz:CL)
  5. EQUITIES: Tesla (xyz:TSLA), Nvidia (xyz:NVDA), MicroStrategy (xyz:MSTR)

Extracted Layers:
  - LEVEL 2 (L2) Aggregated Orderbook: Top 20 bids & asks, unit quantities, USD notionals,
    cumulative depth curves, order counts (n), bid/ask ratio, and live spread (points & bps).
  - LEVEL 3 (L3) Granular Resting Limit Orders: Individual orders mapped directly to verified
    on-chain Ethereum/Hyperliquid wallet addresses (0x...), classified by whale tiers
    ($100K, $1M, $5M), distance from mid %, and direct Hyperdash explorer links.

Outputs to worksheet '17_Orderbook' in Data/DataBase_Trading.xlsx (positioned at Index 25).
"""

from __future__ import annotations

import argparse
import datetime
import json
import logging
import math
import os
import pathlib
import sys
import time
import urllib.request
from typing import Any, Dict, List, Optional, Tuple

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] [%(name)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger("17_Orderbook")

MODULES_DIR = pathlib.Path(__file__).resolve().parent
ROOT_DIR = MODULES_DIR.parents[1]
DEFAULT_EXCEL_PATH = ROOT_DIR / "Data" / "DataBase_Trading.xlsx"
HL_INFO_URL = "https://api.hyperliquid.xyz/info"
HD_GRAPHQL_URL = "https://api.hyperdash.com/graphql"
SHEET_NAME = "17_Orderbook"

# Complete Multi-Asset Roster across Crypto, Indices, Metals, Commodities, and Equities
DEFAULT_ASSET_ROSTER = [
    # (Asset Class, Market Slug, Display Symbol, Min Notional USD)
    ("CRYPTO", "BTC", "BTC", 5000.0),
    ("CRYPTO", "ETH", "ETH", 2500.0),
    ("CRYPTO", "SOL", "SOL", 1000.0),
    ("CRYPTO", "BNB", "BNB", 1000.0),
    ("CRYPTO", "XRP", "XRP", 1000.0),
    ("CRYPTO", "DOGE", "DOGE", 500.0),
    ("CRYPTO", "AVAX", "AVAX", 500.0),
    ("CRYPTO", "SUI", "SUI", 500.0),
    ("INDICES", "xyz:SP500", "SP500", 2000.0),
    ("METALS", "xyz:GOLD", "GOLD", 2000.0),
    ("METALS", "xyz:SILVER", "SILVER", 1000.0),
    ("COMMODITIES", "xyz:CL", "CL_OIL", 1000.0),
    ("EQUITIES", "xyz:TSLA", "TSLA", 2000.0),
    ("EQUITIES", "xyz:NVDA", "NVDA", 2000.0),
    ("EQUITIES", "xyz:MSTR", "MSTR", 2000.0),
]

COLUMNS_DEF = [
    ("sr_no", "Sr.No", 8, "center"),
    ("data_layer", "Data Layer", 14, "center"),
    ("asset_class", "Asset Class", 14, "center"),
    ("asset", "Asset", 12, "center"),
    ("side", "Side", 10, "center"),
    ("price", "Price ($)", 16, "right"),
    ("size_coins", "Size (Units)", 18, "right"),
    ("notional_usd", "Notional (USD)", 18, "right"),
    ("cumulative_depth", "Cumulative Depth ($)", 22, "right"),
    ("orders_count_tier", "Orders (n) / Whale Tier", 22, "center"),
    ("distance_pct", "Distance (%)", 14, "right"),
    ("spread_or_bias", "Spread / Orderflow Bias", 24, "center"),
    ("wallet_address", "Wallet Address", 44, "left"),
    ("explorer_url", "Hyperdash Explorer URL", 45, "left"),
]


def post_json(url: str, payload: dict, timeout: int = 8) -> Any:
    """Post JSON helper with safe timeout, user-agent, and retry."""
    data_bytes = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=data_bytes,
        headers={
            "Content-Type": "application/json",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
        },
    )
    for attempt in range(2):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except Exception as exc:
            if attempt == 1:
                log.warning(f"POST {url} failed on attempt {attempt+1}: {exc}")
            time.sleep(0.3)
    return None


def fetch_l2_book(coin: str) -> Optional[Dict[str, Any]]:
    """Fetch real-time Level 2 orderbook with 20 bids and 20 asks from Hyperliquid."""
    payload = {"type": "l2Book", "coin": coin}
    raw = post_json(HL_INFO_URL, payload, timeout=6)
    if not isinstance(raw, dict) or "levels" not in raw or len(raw["levels"]) < 2:
        log.warning(f"[{coin}] Failed fetching L2 orderbook.")
        return None

    raw_bids = raw["levels"][0]
    raw_asks = raw["levels"][1]
    if not raw_bids or not raw_asks:
        return None

    best_bid = float(raw_bids[0]["px"])
    best_ask = float(raw_asks[0]["px"])
    mid_px = (best_bid + best_ask) / 2.0
    spread = best_ask - best_bid
    spread_bps = (spread / best_bid) * 10000.0

    bids_parsed = []
    cum_bid_usd = 0.0
    for b in raw_bids:
        px = float(b["px"])
        sz = float(b["sz"])
        n_orders = int(b.get("n", 1))
        notional = px * sz
        cum_bid_usd += notional
        dist = ((px - mid_px) / mid_px) * 100.0
        bids_parsed.append({
            "price": px,
            "size_coins": sz,
            "notional_usd": notional,
            "cum_usd": cum_bid_usd,
            "orders_count": n_orders,
            "dist_pct": dist,
        })

    asks_parsed = []
    cum_ask_usd = 0.0
    for a in raw_asks:
        px = float(a["px"])
        sz = float(a["sz"])
        n_orders = int(a.get("n", 1))
        notional = px * sz
        cum_ask_usd += notional
        dist = ((px - mid_px) / mid_px) * 100.0
        asks_parsed.append({
            "price": px,
            "size_coins": sz,
            "notional_usd": notional,
            "cum_usd": cum_ask_usd,
            "orders_count": n_orders,
            "dist_pct": dist,
        })

    total_bid_vol = cum_bid_usd
    total_ask_vol = cum_ask_usd
    total_vol = total_bid_vol + total_ask_vol
    bid_ratio = (total_bid_vol / total_vol * 100.0) if total_vol > 0 else 50.0

    return {
        "coin": coin,
        "best_bid": best_bid,
        "best_ask": best_ask,
        "mid_px": mid_px,
        "spread": spread,
        "spread_bps": spread_bps,
        "bids": bids_parsed,
        "asks": asks_parsed,
        "total_bid_vol": total_bid_vol,
        "total_ask_vol": total_ask_vol,
        "bid_ratio": bid_ratio,
        "ask_ratio": 100.0 - bid_ratio,
    }


def fetch_l3_orders(coin: str, mid_px: float, depth_pct: float = 0.02, min_notional_usd: float = 2000.0) -> List[Dict[str, Any]]:
    """Fetch individual resting limit orders mapped to Ethereum wallet addresses via Hyperdash GraphQL."""
    query = """
    query GetOrderbookSnapshotFiltered($market: String!, $minPrice: Float!, $maxPrice: Float!) {
      orderbookSnapshotFiltered(market: $market, minPrice: $minPrice, maxPrice: $maxPrice) {
        address
        order {
          coin
          side
          limitPx
          sz
        }
      }
    }
    """
    min_px = mid_px * (1.0 - depth_pct)
    max_px = mid_px * (1.0 + depth_pct)

    # For HIP-3 assets (xyz:ASSET), query with the market name
    variables = {"market": coin, "minPrice": float(min_px), "maxPrice": float(max_px)}

    raw = post_json(HD_GRAPHQL_URL, {"query": query, "variables": variables}, timeout=10)
    orders_raw = []
    if isinstance(raw, dict) and "data" in raw and raw["data"] and "orderbookSnapshotFiltered" in raw["data"]:
        orders_raw = raw["data"]["orderbookSnapshotFiltered"] or []

    # If empty and coin has prefix, fallback query with short symbol
    if not orders_raw and ":" in coin:
        short_coin = coin.split(":")[-1]
        raw_fallback = post_json(HD_GRAPHQL_URL, {"query": query, "variables": {"market": short_coin, "minPrice": float(min_px), "maxPrice": float(max_px)}}, timeout=8)
        if isinstance(raw_fallback, dict) and "data" in raw_fallback and raw_fallback["data"] and "orderbookSnapshotFiltered" in raw_fallback["data"]:
            orders_raw = raw_fallback["data"]["orderbookSnapshotFiltered"] or []

    parsed_orders = []
    for item in orders_raw:
        sub = item.get("order", {})
        addr = item.get("address", "")
        side_code = sub.get("side", "")
        if side_code not in ("B", "A"):
            continue

        try:
            px = float(sub.get("limitPx", 0.0))
            sz = float(sub.get("sz", 0.0))
        except (ValueError, TypeError):
            continue

        notional = px * sz
        if notional < min_notional_usd or px <= 0 or sz <= 0:
            continue

        side = "BUY" if side_code == "B" else "SELL"
        dist = ((px - mid_px) / mid_px) * 100.0

        # Institutional Whale Tiering
        if notional >= 5_000_000.0:
            tier = "MEGA WHALE (>= $5M)"
            bias = "INSTITUTIONAL_WALL"
        elif notional >= 1_000_000.0:
            tier = "WHALE (>= $1M)"
            bias = "HEAVY_SUPPORT" if side == "BUY" else "HEAVY_RESISTANCE"
        elif notional >= 100_000.0:
            tier = "LARGE (>= $100K)"
            bias = "LIQUIDITY_SHELF"
        elif notional >= 25_000.0:
            tier = "MEDIUM (>= $25K)"
            bias = "ACTIVE_LIMIT"
        else:
            tier = "STANDARD"
            bias = "RESTING_ORDER"

        explorer_url = f"https://hyperdash.com/address/{addr}" if addr.startswith("0x") else "-"

        parsed_orders.append({
            "asset": coin.split(":")[-1],
            "side": side,
            "price": px,
            "size_coins": sz,
            "notional_usd": notional,
            "tier": tier,
            "dist_pct": dist,
            "bias": bias,
            "wallet_address": addr,
            "explorer_url": explorer_url,
        })

    # Sort descending by notional value to feature highest impact whale orders first
    parsed_orders.sort(key=lambda o: o["notional_usd"], reverse=True)
    return parsed_orders


def format_px(px: Optional[float]) -> str:
    if px is None or not math.isfinite(px):
        return "-"
    return f"${px:,.2f}" if px >= 1.0 else f"${px:,.4f}"


def format_usd(amt: Optional[float]) -> str:
    if amt is None or not math.isfinite(amt):
        return "-"
    return f"${amt:,.2f}"


def build_orderbook_dataset(
    asset_roster: List[Tuple[str, str, str, float]] = DEFAULT_ASSET_ROSTER,
    max_l3_per_symbol: int = 350
) -> List[Dict[str, Any]]:
    """Builds unified multi-asset L2 & L3 dataset across all asset classes."""
    master_records: List[Dict[str, Any]] = []
    sr_no = 1

    for asset_class, market_slug, display_sym, min_usd in asset_roster:
        log.info(f"Harvesting [{asset_class}] {display_sym} ({market_slug})...")
        l2 = fetch_l2_book(market_slug)
        if not l2:
            continue

        mid_px = l2["mid_px"]
        spread_text = f"Spread: ${l2['spread']:,.2f} ({l2['spread_bps']:.2f} bps)"
        imbalance_text = f"B/A: {l2['bid_ratio']:.1f}% / {l2['ask_ratio']:.1f}%"
        asset_url = f"https://hyperdash.com/asset/{display_sym}"

        # --- SECTION 1: L2 AGGREGATED LADDER ---
        # 1. Asks (Top 20 descending: high price to best ask)
        for a in reversed(l2["asks"]):
            master_records.append({
                "sr_no": sr_no,
                "data_layer": "L2_AGGREGATED",
                "asset_class": asset_class,
                "asset": display_sym,
                "side": "ASK",
                "price": format_px(a["price"]),
                "size_coins": f"{a['size_coins']:,.4f} {display_sym}",
                "notional_usd": format_usd(a["notional_usd"]),
                "cumulative_depth": format_usd(a["cum_usd"]),
                "orders_count_tier": f"{a['orders_count']} orders",
                "distance_pct": f"{a['dist_pct']:+.2f}%",
                "spread_or_bias": imbalance_text,
                "wallet_address": "-",
                "explorer_url": asset_url,
            })
            sr_no += 1

        # 2. Spread Divider Row
        master_records.append({
            "sr_no": sr_no,
            "data_layer": "L2_SPREAD",
            "asset_class": asset_class,
            "asset": display_sym,
            "side": "MID_SPREAD",
            "price": format_px(mid_px),
            "size_coins": "-",
            "notional_usd": "-",
            "cumulative_depth": "-",
            "orders_count_tier": f"Mid: ${mid_px:,.2f}",
            "distance_pct": "0.00%",
            "spread_or_bias": spread_text,
            "wallet_address": "-",
            "explorer_url": asset_url,
        })
        sr_no += 1

        # 3. Bids (Top 20 descending: best bid to low price)
        for b in l2["bids"]:
            master_records.append({
                "sr_no": sr_no,
                "data_layer": "L2_AGGREGATED",
                "asset_class": asset_class,
                "asset": display_sym,
                "side": "BID",
                "price": format_px(b["price"]),
                "size_coins": f"{b['size_coins']:,.4f} {display_sym}",
                "notional_usd": format_usd(b["notional_usd"]),
                "cumulative_depth": format_usd(b["cum_usd"]),
                "orders_count_tier": f"{b['orders_count']} orders",
                "distance_pct": f"{b['dist_pct']:+.2f}%",
                "spread_or_bias": imbalance_text,
                "wallet_address": "-",
                "explorer_url": asset_url,
            })
            sr_no += 1

        # --- SECTION 2: L3 GRANULAR RESTING ORDERS WITH WALLET ATTRIBUTION ---
        time.sleep(0.25)  # Rate-limit guard between GraphQL calls
        l3_orders = fetch_l3_orders(market_slug, mid_px=mid_px, depth_pct=0.02, min_notional_usd=min_usd)
        log.info(f"[{display_sym}] Captured {len(l3_orders)} granular L3 orders (>= ${min_usd:,.0f}).")

        capped_l3 = l3_orders[:max_l3_per_symbol]
        cum_wall = 0.0
        for o in capped_l3:
            cum_wall += o["notional_usd"]
            master_records.append({
                "sr_no": sr_no,
                "data_layer": "L3_RESTING_ORDER",
                "asset_class": asset_class,
                "asset": display_sym,
                "side": o["side"],
                "price": format_px(o["price"]),
                "size_coins": f"{o['size_coins']:,.4f} {display_sym}",
                "notional_usd": format_usd(o["notional_usd"]),
                "cumulative_depth": format_usd(cum_wall),
                "orders_count_tier": o["tier"],
                "distance_pct": f"{o['dist_pct']:+.2f}%",
                "spread_or_bias": o["bias"],
                "wallet_address": o["wallet_address"],
                "explorer_url": o["explorer_url"],
            })
            sr_no += 1

    return master_records


def dump_to_excel(records: List[Dict[str, Any]], excel_path: pathlib.Path = DEFAULT_EXCEL_PATH) -> bool:
    """Dumps structured multi-asset L2 & L3 Orderbook into worksheet '17_Orderbook'."""
    if openpyxl is None:
        log.error("openpyxl library not installed.")
        return False

    if not excel_path.exists():
        log.error(f"Excel workbook not found at: {excel_path}")
        return False

    if not records:
        log.error("0 records to dump. Aborting to protect workbook integrity.")
        return False

    wb = openpyxl.load_workbook(excel_path)

    # Position at Index 25 (immediately after 16_Hyperdash_telegram_Intel at Index 24)
    if SHEET_NAME in wb.sheetnames:
        idx = wb.sheetnames.index(SHEET_NAME)
        wb.remove(wb[SHEET_NAME])
        ws = wb.create_sheet(title=SHEET_NAME, index=idx)
    else:
        if "16_Hyperdash_telegram_Intel" in wb.sheetnames:
            target_idx = wb.sheetnames.index("16_Hyperdash_telegram_Intel") + 1
        elif "15_Hyperdash_Stops_Map" in wb.sheetnames:
            target_idx = wb.sheetnames.index("15_Hyperdash_Stops_Map") + 2
        else:
            target_idx = 25 if len(wb.sheetnames) >= 25 else len(wb.sheetnames)
        ws = wb.create_sheet(title=SHEET_NAME, index=target_idx)

    final_idx = wb.sheetnames.index(SHEET_NAME)
    log.info(f"Target worksheet '{SHEET_NAME}' positioned at index {final_idx}")

    # Typography & Colors
    hdr_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    base_font = Font(name="Segoe UI", size=9, color="000000")
    bold_font = Font(name="Segoe UI", size=9, bold=True, color="000000")
    mono_font = Font(name="Consolas", size=9, color="334155")
    green_font = Font(name="Segoe UI", size=9, bold=True, color="00A86B")
    red_font = Font(name="Segoe UI", size=9, bold=True, color="E53935")
    amber_font = Font(name="Segoe UI", size=9, bold=True, color="D97706")
    purple_font = Font(name="Segoe UI", size=9, bold=True, color="7C3AED")

    spread_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    whale_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    mega_fill = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")

    # Category Badges Font & Fill
    badge_colors = {
        "CRYPTO": ("0284C7", "E0F2FE"),       # Sky Blue
        "INDICES": ("1E293B", "F1F5F9"),      # Slate Navy
        "METALS": ("B45309", "FEF3C7"),       # Amber/Gold
        "COMMODITIES": ("3F6212", "ECFCCB"),  # Olive Green
        "EQUITIES": ("6D28D9", "EDE9FE"),     # Purple
    }

    thin_border_side = Side(border_style="thin", color="E2E8F0")
    row_border = Border(
        left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side
    )

    # 1. Header Row
    ws.row_dimensions[1].height = 26
    for col_idx, (col_key, header_title, width, align) in enumerate(COLUMNS_DEF, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_title)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.border = row_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 2. Data Rows
    for row_idx, r in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 20
        dlayer = r.get("data_layer", "")
        side_val = r.get("side", "")
        tier = r.get("orders_count_tier", "")
        aclass = r.get("asset_class", "CRYPTO")

        is_spread_row = dlayer == "L2_SPREAD"
        is_mega_whale = "MEGA" in tier
        is_whale = "WHALE" in tier and not is_mega_whale

        for col_idx, (col_key, _, _, align) in enumerate(COLUMNS_DEF, start=1):
            val = r.get(col_key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(horizontal=align, vertical="center")
            cell.border = row_border

            # Background highlighting
            if is_spread_row:
                cell.fill = spread_fill
            elif is_mega_whale and col_key in ("notional_usd", "orders_count_tier", "wallet_address"):
                cell.fill = mega_fill
            elif is_whale and col_key in ("notional_usd", "orders_count_tier", "wallet_address"):
                cell.fill = whale_fill
            elif col_key == "asset_class" and aclass in badge_colors:
                _, bg_col = badge_colors[aclass]
                cell.fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")

            # Typography formatting
            if is_spread_row:
                cell.font = bold_font
            elif col_key in ("sr_no", "asset"):
                cell.font = bold_font
            elif col_key == "asset_class" and aclass in badge_colors:
                fg_col, _ = badge_colors[aclass]
                cell.font = Font(name="Segoe UI", size=9, bold=True, color=fg_col)
            elif col_key == "data_layer":
                cell.font = bold_font if dlayer == "L2_AGGREGATED" else base_font
            elif col_key == "side":
                cell.font = green_font if side_val in ("BID", "BUY") else (red_font if side_val in ("ASK", "SELL") else bold_font)
            elif col_key == "notional_usd":
                cell.font = purple_font if is_mega_whale else (amber_font if is_whale else bold_font)
            elif col_key == "orders_count_tier":
                cell.font = purple_font if is_mega_whale else (amber_font if is_whale else base_font)
            elif col_key == "distance_pct":
                cell.font = green_font if str(val).startswith("+") else (red_font if str(val).startswith("-") else base_font)
            elif col_key == "wallet_address":
                cell.font = mono_font
            else:
                cell.font = base_font

    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    wb.save(excel_path)
    log.info(f"Successfully saved {len(records)} rows to '{SHEET_NAME}' at {excel_path}")
    return True


def main():
    parser = argparse.ArgumentParser(description="17_Orderbook Multi-Asset Orderbook Harvester")
    parser.add_argument("--categories", nargs="+", default=["CRYPTO", "INDICES", "METALS", "COMMODITIES", "EQUITIES"], help="Asset classes to harvest")
    parser.add_argument("--max-l3", type=int, default=350, help="Max L3 orders per symbol")
    parser.add_argument("--excel-path", type=pathlib.Path, default=DEFAULT_EXCEL_PATH, help="Path to DataBase_Trading.xlsx")
    args = parser.parse_args()

    selected_categories = {c.upper() for c in args.categories}
    roster = [item for item in DEFAULT_ASSET_ROSTER if item[0] in selected_categories]

    log.info(f"Starting Multi-Asset Orderbook harvest for {len(roster)} markets across {selected_categories}")
    records = build_orderbook_dataset(roster, max_l3_per_symbol=args.max_l3)

    if not records:
        log.error("Harvesting yielded 0 records. Excel dump aborted.")
        sys.exit(1)

    success = dump_to_excel(records, excel_path=args.excel_path)
    if not success:
        sys.exit(1)

    log.info("17_Orderbook pipeline executed successfully.")


if __name__ == "__main__":
    main()
