#!/usr/bin/env python3
"""
Data/Excel_Modules/22_Forced_Liquidations_MTF.py
=================================================
Excel Exporter for Sheet '22_Forced_Liquidations_MTF' in Data/DataBase_Trading.xlsx.
Conforms strictly to xlsx-official formatting guidelines:
  - Row 1: Banner with Title, Timestamp (IST + UTC), Row Count, Status
  - Row 2: Formatted Headers with frozen panes at A3
  - Data Rows: Formatted numbers ($ currency, commas, decimals)
  - Color Fills: Soft Coral for Long Liq, Soft Green for Short Liq, Amber for Climax
  - Dual Timestamps: Explicit IST (UTC+5:30) alongside UTC
"""

from __future__ import annotations

import datetime
import os
import pathlib
import sys
import tempfile
import time
from typing import Dict, List, Optional

try:
    from logger_config import stamp_sheet_completion_ist
except ImportError:
    pass

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

try:
    from filelock import FileLock
except ImportError:
    FileLock = None

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "Data"
DEFAULT_EXCEL_PATH = DATA_DIR / "DataBase_Trading.xlsx"

IST_TIMEZONE = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
UTC_TIMEZONE = datetime.timezone.utc
SHEET_NAME = "22_Forced_Liquidations_MTF"


def dump_forced_liquidations_to_excel(
    candles: List[Dict],
    target_path: Optional[pathlib.Path] = None
) -> bool:
    """Writes aggregated multi-timeframe liquidation candles into target Excel workbook."""
    if openpyxl is None:
        print("[ERROR] openpyxl not installed.")
        return False

    excel_path = target_path or DEFAULT_EXCEL_PATH
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    if not excel_path.exists():
        wb = openpyxl.Workbook()
        wb.active.title = SHEET_NAME
        ws = wb.active
    else:
        try:
            wb = openpyxl.load_workbook(excel_path)
            if SHEET_NAME in wb.sheetnames:
                ws = wb[SHEET_NAME]
                ws.delete_rows(1, ws.max_row + 1)
                # Ensure Sheet 22 is positioned at the very end
                if ws in wb._sheets and wb._sheets[-1] != ws:
                    wb._sheets.remove(ws)
                    wb._sheets.append(ws)
            else:
                ws = wb.create_sheet(title=SHEET_NAME)
        except Exception as e:
            print(f"[ERROR] Failed to load workbook: {e}")
            return False

    # Styling Palettes (xlsx-official)
    font_banner = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_banner = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")  # Navy

    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")  # Slate

    font_data = Font(name="Calibri", size=9)
    font_bold = Font(name="Calibri", size=9, bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    fill_short_liq = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")  # Mint
    fill_long_liq = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")   # Rose
    fill_climax = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")     # Amber highlight

    now_utc = datetime.datetime.now(UTC_TIMEZONE)
    now_ist = now_utc.astimezone(IST_TIMEZONE)
    ist_str = now_ist.strftime("%Y-%m-%d %H:%M:%S IST")
    utc_str = now_utc.strftime("%Y-%m-%d %H:%M:%S UTC")

    # 1. Row 1 Banner
    banner_text = f"{SHEET_NAME} | LAST UPDATED: {ist_str} ({utc_str}) | ROWS: {len(candles)} | STATUS: LIVE"
    ws.merge_cells("A1:M1")
    banner_cell = ws["A1"]
    banner_cell.value = banner_text
    banner_cell.font = font_banner
    banner_cell.fill = fill_banner
    banner_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    # 2. Row 2 Column Headers
    headers = [
        ("#", 6),
        ("Asset / Pair", 14),
        ("TF", 8),
        ("Candle Start (IST)", 22),
        ("Candle End (IST)", 22),
        ("Candle Start (UTC)", 22),
        ("Long Liq ($)", 16),
        ("Short Liq ($)", 16),
        ("Total Liq ($)", 16),
        ("Net Liq Delta ($)", 18),
        ("Orders", 10),
        ("Dominant Flow", 18),
        ("Climax Status", 16),
    ]

    ws.row_dimensions[2].height = 24
    for col_idx, (h_title, width) in enumerate(headers, start=1):
        c = ws.cell(row=2, column=col_idx, value=h_title)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 3. Data Rows
    row_num = 3
    for idx, c_data in enumerate(candles, start=1):
        ws.row_dimensions[row_num].height = 20
        total_liq = float(c_data.get("total_liq_usd", 0.0))
        long_liq = float(c_data.get("long_liq_usd", 0.0))
        short_liq = float(c_data.get("short_liq_usd", 0.0))
        net_delta = float(c_data.get("net_liq_delta_usd", 0.0))
        orders = int(c_data.get("orders_count", 0))
        flow = str(c_data.get("dominant_side", "BALANCED"))
        is_climax = total_liq >= 500000.0 or orders >= 25

        # Format values
        row_vals = [
            (idx, align_center, font_data, None),
            (c_data.get("symbol", ""), align_left, font_bold, None),
            (c_data.get("timeframe", ""), align_center, font_bold, None),
            (c_data.get("candle_start_ist", ""), align_center, font_data, None),
            (c_data.get("candle_end_ist", ""), align_center, font_data, None),
            (c_data.get("candle_start_utc", ""), align_center, font_data, None),
            (f"${long_liq:,.2f}", align_right, font_data, fill_long_liq if long_liq > 0 else None),
            (f"${short_liq:,.2f}", align_right, font_data, fill_short_liq if short_liq > 0 else None),
            (f"${total_liq:,.2f}", align_right, font_bold, fill_climax if is_climax else None),
            (f"${net_delta:+,.2f}", align_right, font_bold, fill_short_liq if net_delta > 0 else fill_long_liq if net_delta < 0 else None),
            (orders, align_right, font_data, None),
            (flow, align_center, font_bold, fill_short_liq if flow == "SHORTS_PURGED" else fill_long_liq if flow == "LONGS_FLUSHED" else None),
            ("⚠️ CLIMAX SPIKE" if is_climax else "NORMAL", align_center, font_bold, fill_climax if is_climax else None),
        ]

        for col_idx, (val, alignment, font, fill) in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.alignment = alignment
            cell.font = font
            cell.border = thin_border
            if fill:
                cell.fill = fill

        row_num += 1

    # Freeze panes at A3
    ws.freeze_panes = "A3"
    ws.views.sheetView[0].showGridLines = True

    # Save workbook
    try:
        wb.save(excel_path)
        wb.close()
        return True
    except Exception as e:
        print(f"[ERROR] Failed to save workbook directly: {e}")
        try:
            wb.close()
        except Exception:
            pass
        return False


if __name__ == "__main__":
    # Test script run
    print("Testing 22_Forced_Liquidations_MTF module...")
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from Terminal.Data_Factory.forced_liquidation_daemon import ForcedLiquidationAggregator
    agg = ForcedLiquidationAggregator()
    agg.bootstrap_from_excel_live_liq(DEFAULT_EXCEL_PATH)
    candles = agg.get_all_candles_sorted(limit=500)
    ok = dump_forced_liquidations_to_excel(candles)
    print(f"Dump complete: success={ok}, candles={len(candles)}")
