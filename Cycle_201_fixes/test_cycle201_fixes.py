import os
import sys

sys.path.insert(0, os.path.dirname(__file__))

import openpyxl  # noqa: E402
import pytest  # noqa: E402

from fail_closed import (  # noqa: E402
    NA, num, cell, spread_bps, spread_gate, stressed_equity, admissible,
    max_slots, four_slot_threshold, bar_fresh,
)
from atomic_xlsx import atomic_save, check_minimums, cache_sum_formulas  # noqa: E402
import atomic_xlsx  # noqa: E402


# ---- F1: numeric parsing, strict
@pytest.mark.parametrize("v", [None, "", "-", "nan", "inf", "-inf", "1e400",
                               "abc", True, False, float("nan"), float("inf"),
                               "1,5", "1,234,5", "1_000", "12,34", " ", "0x1A"])
def test_num_rejects_unknown_or_ambiguous(v):
    assert num(v) is None


@pytest.mark.parametrize("v,expected", [
    ("1,234.5", 1234.5), ("1,058,753.36", 1058753.36), ("0", 0.0), (0, 0.0),
    ("-1.5", -1.5), (" 81722 ", 81722.0), (".5", 0.5), ("1e3", 1000.0), (7, 7.0),
])
def test_num_accepts_plain_and_grouped(v, expected):
    assert num(v) == pytest.approx(expected)


def test_cell_renders_na_not_zero():
    assert cell(None) == NA
    assert cell("1,5") == NA           # ambiguous -> N/A, not 15
    assert cell(1.23456, "{:.2f}") == 1.23


# ---- F5: spread, locked and stale quotes fail closed
def test_spread_matches_live_btc_snapshot():
    assert round(spread_bps(81722.0, 81738.0, 0.1), 2) == 1.96


def test_spread_locked_quote_is_none_not_zero():
    assert spread_bps(100.0, 100.0, 0.1) is None      # was 0.0 in Cycle 200
    assert spread_gate(spread_bps(100.0, 100.0, 0.1), "CRYPTO") is False


def test_spread_invalid_or_stale_is_none():
    assert spread_bps(None, 81738.0, 0.1) is None
    assert spread_bps(81738.0, 81722.0, 0.1) is None   # crossed
    assert spread_bps(81722.0, 81738.0, 6.0) is None   # stale > 5 s
    assert spread_bps(81722.0, 81738.0, -1.0) is None  # negative age
    assert spread_bps(0.0, 1.0, 0.1) is None


def test_spread_gate_limits_and_fail_closed():
    assert spread_gate(25.0, "CRYPTO") is True
    assert spread_gate(25.06, "CRYPTO") is False       # live SOL fails
    assert spread_gate(23.17, "CRYPTO") is True        # live XRP passes
    assert spread_gate(5.0, "TRADFI") is True
    assert spread_gate(6.76, "TRADFI") is False        # live EURUSD fails
    assert spread_gate(None, "CRYPTO") is False
    assert spread_gate(0.0, "UNKNOWN") is False


# ---- F6: capacity
def test_stressed_equity_formula():
    assert stressed_equity(4863.74, [25.0, 25.0]) == pytest.approx(4797.24)
    assert stressed_equity(4863.74, [25.0, None]) is None
    assert stressed_equity(4863.74, [-5.0]) is None


def test_max_slots_matches_remediation_claims():
    assert max_slots(4863.74, 25.0) == 2
    assert max_slots(4863.74, 15.0) == 3


def test_max_slots_accepts_numeric_string_risk():
    assert max_slots(4863.74, "15") == 3              # was TypeError in Cycle 200
    assert max_slots(4863.74, "15,0") == 0           # ambiguous string -> 0 slots
    assert max_slots(4863.74, "abc") == 0


def test_four_slot_threshold_4878_and_unknown_risk():
    assert four_slot_threshold(15.0) == pytest.approx(4878.0)
    assert four_slot_threshold("15") == pytest.approx(4878.0)
    assert four_slot_threshold(None) is None
    assert max_slots(4878.0, 15.0) == 4
    assert max_slots(4877.99, 15.0) == 3


def test_admissible_fail_closed():
    assert admissible(4863.74, [], 25.0) is True
    assert admissible(4863.74, [25.0], 25.0) is True
    assert admissible(4863.74, [25.0, 25.0], 25.0) is False
    assert admissible(4863.74, [None], 15.0) is False
    assert admissible(4863.74, [15.0] * 4, 15.0) is False
    assert admissible(4863.74, [], None) is False


# ---- freshness
def test_bar_freshness():
    now = 1_000_000.0
    assert bar_fresh(now - 1000, now) is True          # open-age in [900,1800)
    assert bar_fresh(now - 2000, now) is False         # one bar stale
    assert bar_fresh(now + 60, now) is False           # future-dated
    assert bar_fresh(None, now) is False
    assert bar_fresh(now - 60, now, session_open=False) is False


# ---- F2: formula cache, blank = N/A by default
def _sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["F3"], ws["G3"], ws["H3"] = 1058753.36, 743422.24, 1077880.33
    ws["I3"] = "=F3+G3+H3"
    ws["F4"], ws["G4"], ws["H4"] = 1.0, None, 2.0       # blank OKX leg
    ws["I4"] = "=F4+G4+H4"
    ws["F5"], ws["G5"], ws["H5"] = "-", 1.0, 2.0        # dash text
    ws["I5"] = "=F5+G5+H5"
    ws["F6"], ws["G6"], ws["H6"] = "1,058.5", 1.0, 2.0  # numeric text with grouping
    ws["I6"] = "=F6+G6+H6"
    ws["I7"] = "=SUM(F7:H7)"                            # unsupported -> skipped
    return wb, ws


def test_formula_cache_blank_is_na_by_default():
    _, ws = _sheet()
    stats = cache_sum_formulas(ws)
    assert ws["I3"].value == round(1058753.36 + 743422.24 + 1077880.33, 2)
    assert ws["I4"].value == NA                         # was 3.0 in Cycle 200
    assert ws["I5"].value == NA
    assert ws["I6"].value == round(1058.5 + 1.0 + 2.0, 2)
    assert ws["I7"].value == "=SUM(F7:H7)"
    assert stats == {"cached": 2, "na": 2, "skipped": 1}


def test_formula_cache_blank_zero_is_explicit_opt_in():
    _, ws = _sheet()
    cache_sum_formulas(ws, blank_policy="zero")
    assert ws["I4"].value == 3.0


def test_formula_cache_rejects_bad_policy():
    _, ws = _sheet()
    with pytest.raises(ValueError):
        cache_sum_formulas(ws, blank_policy="guess")


# ---- F3: atomic save, verify before replace, lock handling, minimums
def test_atomic_save_creates_bak_and_cleans_tmp(tmp_path):
    wb, _ = _sheet()
    p = tmp_path / "DataBase_Trading.xlsx"
    atomic_save(wb, p, expected_sheets=1)
    assert p.exists()
    assert not (tmp_path / "DataBase_Trading.tmp.xlsx").exists()
    atomic_save(wb, p, expected_sheets=1)
    assert (tmp_path / "DataBase_Trading.bak.xlsx").exists()


def test_failed_verify_leaves_original_untouched(tmp_path):
    wb, _ = _sheet()
    p = tmp_path / "DataBase_Trading.xlsx"
    atomic_save(wb, p)
    before = p.read_bytes()
    with pytest.raises(RuntimeError):
        atomic_save(wb, p, verify=lambda t: False)
    assert p.read_bytes() == before


def test_wrong_sheet_count_leaves_original_untouched(tmp_path):
    wb, _ = _sheet()
    p = tmp_path / "DataBase_Trading.xlsx"
    atomic_save(wb, p)
    before = p.read_bytes()
    with pytest.raises(RuntimeError, match="sheet count"):
        atomic_save(wb, p, expected_sheets=40)
    assert p.read_bytes() == before


def test_locked_target_raises_and_keeps_original(tmp_path, monkeypatch):
    wb, _ = _sheet()
    p = tmp_path / "DataBase_Trading.xlsx"
    atomic_save(wb, p)
    before = p.read_bytes()

    def locked(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(atomic_xlsx.os, "replace", locked)
    with pytest.raises(RuntimeError, match="locked"):
        atomic_save(wb, p)
    assert p.read_bytes() == before
    assert (tmp_path / "DataBase_Trading.tmp.xlsx").exists()   # new copy kept


def test_check_minimums_missing_key_fails_closed():
    with pytest.raises(RuntimeError):
        check_minimums({}, {"arkham": 50})
    with pytest.raises(RuntimeError):
        check_minimums({"arkham": 0}, {"arkham": 50})
    check_minimums({"arkham": 138}, {"arkham": 50})
