import math
import os
import sys
import tempfile

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

import pytest  # noqa: E402
import openpyxl  # noqa: E402

try:
    from Cycle_200_fixes.fail_closed import (  # noqa: E402
        NA, num, cell, spread_bps, spread_gate, stressed_equity, admissible,
        max_slots, four_slot_threshold, bar_fresh,
    )
    from Cycle_200_fixes.atomic_xlsx import atomic_save, check_minimums, cache_sum_formulas  # noqa: E402
except ImportError:
    from fail_closed import (  # noqa: E402
        NA, num, cell, spread_bps, spread_gate, stressed_equity, admissible,
        max_slots, four_slot_threshold, bar_fresh,
    )
    from atomic_xlsx import atomic_save, check_minimums, cache_sum_formulas  # noqa: E402


# ---------------- numeric parsing: no fake defaults
@pytest.mark.parametrize("v", [None, "", "-", "nan", "inf", "abc", True, float("nan")])
def test_num_rejects_unknown(v):
    assert num(v) is None


def test_num_accepts_numbers_and_thousands_separators():
    assert num("1,234.5") == 1234.5
    assert num(0) == 0.0          # real zero is kept, not confused with missing
    assert num("0") == 0.0


def test_cell_renders_na_not_zero():
    assert cell(None) == NA
    assert cell("-") == NA
    assert cell(1.23456, "{:.2f}") == 1.23


# ---------------- spread: one definition, fail closed
def test_spread_matches_live_btc_snapshot():
    # live snapshot BTC: bid 81722, ask 81738, spread_bps 1.96
    assert round(spread_bps(81722.0, 81738.0, 0.1), 2) == 1.96


def test_spread_invalid_or_stale_is_none():
    assert spread_bps(None, 81738.0, 0.1) is None
    assert spread_bps(81738.0, 81722.0, 0.1) is None       # crossed
    assert spread_bps(81722.0, 81738.0, 6.0) is None        # stale > 5 s
    assert spread_bps(0.0, 1.0, 0.1) is None                # non-positive bid


def test_spread_gate_limits_and_fail_closed():
    assert spread_gate(25.0, "CRYPTO") is True
    assert spread_gate(25.06, "CRYPTO") is False            # live SOL fails
    assert spread_gate(23.17, "CRYPTO") is True             # live XRP passes
    assert spread_gate(5.0, "TRADFI") is True
    assert spread_gate(6.76, "TRADFI") is False             # live EURUSD fails
    assert spread_gate(None, "CRYPTO") is False             # missing never passes
    assert spread_gate(0.0, "UNKNOWN") is False             # unknown category


# ---------------- capacity: arithmetic from Cycle 200 remediation, verified
def test_stressed_equity_formula():
    assert stressed_equity(4863.74, [25.0, 25.0]) == pytest.approx(4797.24)
    assert stressed_equity(4863.74, []) == 4863.74
    assert stressed_equity(4863.74, [25.0, None]) is None


def test_max_slots_matches_remediation_claims():
    assert max_slots(4863.74, 25.0) == 2     # floor(68.74/33.25)
    assert max_slots(4863.74, 15.0) == 3     # floor(68.74/20.75)


def test_four_slot_threshold_is_4878_at_15_usd():
    assert four_slot_threshold(15.0) == pytest.approx(4878.0)
    assert max_slots(4878.0, 15.0) == 4
    assert max_slots(4877.99, 15.0) == 3


def test_max_slots_unknown_or_negative_is_zero():
    assert max_slots(None, 15.0) == 0
    assert max_slots(4700.0, 15.0) == 0
    assert max_slots(4863.74, 0.0) == 0


def test_admissible_fail_closed_on_unknown_open_risk():
    assert admissible(4863.74, [], 25.0) is True
    assert admissible(4863.74, [25.0], 25.0) is True        # 4830.49
    assert admissible(4863.74, [25.0, 25.0], 25.0) is False # 4763.99 < 4795
    assert admissible(4863.74, [None], 15.0) is False
    assert admissible(4863.74, [15.0] * 4, 15.0) is False   # 5th slot over cap


# ---------------- freshness
def test_bar_freshness_crypto():
    now = 1_000_000.0
    assert bar_fresh(now - 600, now) is True        # 10 min old, 15m bars
    assert bar_fresh(now - 2000, now) is False      # > 2 bars
    assert bar_fresh(None, now) is False
    assert bar_fresh(now - 60, now, session_open=False) is False


# ---------------- workbook: formula cache mirrors Excel
def _sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["F3"], ws["G3"], ws["H3"] = 1058753.36, 743422.24, 1077880.33
    ws["I3"] = "=F3+G3+H3"
    ws["F4"], ws["G4"], ws["H4"] = 1.0, None, 2.0          # blank -> 0
    ws["I4"] = "=F4+G4+H4"
    ws["F5"], ws["G5"], ws["H5"] = "-", 1.0, 2.0           # text -> N/A
    ws["I5"] = "=F5+G5+H5"
    ws["I6"] = "=SUM(F6:H6)"                               # unsupported -> skipped
    return wb, ws


def test_formula_cache_values_and_na():
    wb, ws = _sheet()
    stats = cache_sum_formulas(ws, blank_policy="zero")
    assert ws["I3"].value == round(1058753.36 + 743422.24 + 1077880.33, 2)
    assert ws["I4"].value == 3.0
    assert ws["I5"].value == NA
    assert ws["I6"].value == "=SUM(F6:H6)"
    assert stats == {"cached": 2, "na": 1, "skipped": 1}


def test_atomic_save_and_backup_and_minimums():
    wb, _ = _sheet()
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "DataBase_Trading.xlsx")
        atomic_save(wb, p)
        assert os.path.exists(p) and not os.path.exists(p.replace(".xlsx", ".tmp.xlsx"))
        atomic_save(wb, p)
        assert os.path.exists(p.replace(".xlsx", ".bak.xlsx"))
    with pytest.raises(RuntimeError):
        check_minimums({"arkham": 0}, {"arkham": 50})
    check_minimums({"arkham": 138}, {"arkham": 50})   # passes
