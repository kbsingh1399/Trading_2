# Tests/Test_Forced_Liquidation_Excel.py
import pathlib
import openpyxl
import pytest

from Data.Excel_Modules._22_Forced_Liquidations_MTF import dump_forced_liquidations_to_excel

def test_forced_liquidation_sheet_structure(tmp_path):
    excel_path = tmp_path / "Test_Trading.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "00_MT5_Account_Telemetry"
    wb.save(excel_path)

    sample_data = [
        {
            "symbol": "BTCUSDT",
            "timeframe": "15m",
            "candle_start_utc": "2026-10-25 00:00:00 UTC",
            "candle_start_ist": "2026-10-25 05:30:00 IST",
            "candle_end_ist": "2026-10-25 05:45:00 IST",
            "long_liq_usd": 250000.0,
            "short_liq_usd": 1500000.0,
            "total_liq_usd": 1750000.0,
            "net_liq_delta_usd": 1250000.0,
            "orders_count": 18,
            "dominant_side": "SHORTS_PURGED",
        }
    ]

    success = dump_forced_liquidations_to_excel(sample_data, target_path=excel_path)
    assert success is True

    wb_read = openpyxl.load_workbook(excel_path, read_only=True)
    assert "22_Forced_Liquidations_MTF" in wb_read.sheetnames
    ws = wb_read["22_Forced_Liquidations_MTF"]
    banner = ws["A1"].value
    assert "22_Forced_Liquidations_MTF" in banner
    assert "IST" in banner
    assert ws["A2"].value == "#"
    assert ws["D2"].value == "Candle Start (IST)"
    assert ws["E2"].value == "Candle End (IST)"
