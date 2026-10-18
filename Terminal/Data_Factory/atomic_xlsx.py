"""Terminal/Data_Factory/atomic_xlsx.py
======================================
Canonical atomic workbook save, row-count guard, and formula evaluation cache.
"""
from __future__ import annotations

import os
import re
import shutil
from pathlib import Path

try:
    from Terminal.risk.fail_closed import NA, num
except ImportError:
    try:
        from ..risk.fail_closed import NA, num
    except ImportError:
        from fail_closed import NA, num

import openpyxl

_SUM_FORMULA = re.compile(
    r"^=\s*([A-Z]+)(\d+)\s*\+\s*([A-Z]+)(\d+)\s*\+\s*([A-Z]+)(\d+)\s*$"
)


def _default_verify(path: Path) -> bool:
    """Re-open the written temp file. Raises if it is not a readable workbook."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return len(wb.sheetnames) > 0
    finally:
        wb.close()


def atomic_save(wb, path, verify=None, expected_sheets: int | None = None) -> None:
    """Write temp -> fsync -> verify temp -> keep .bak -> os.replace.

    * The live file is replaced only after the temp file re-opens and passes
      verification (sheet count == expected_sheets when given).
    * If the target is locked (e.g. open in Excel) os.replace raises
      PermissionError; we raise RuntimeError and the original is left intact.
    """
    path = Path(path)
    tmp = path.with_name(path.stem + ".tmp" + path.suffix)
    bak = path.with_name(path.stem + ".bak" + path.suffix)

    wb.save(tmp)
    with open(tmp, "a+b") as fh:
        os.fsync(fh.fileno())

    ok = (verify(tmp) if verify is not None else _default_verify(tmp))
    if not ok:
        raise RuntimeError(f"verification failed for {tmp}; original left intact")
    if expected_sheets is not None:
        _wb_check = openpyxl.load_workbook(tmp, read_only=True)
        try:
            n = len(_wb_check.sheetnames)
        finally:
            _wb_check.close()
        if n != expected_sheets:
            raise RuntimeError(f"sheet count {n} != expected {expected_sheets}; original left intact")

    if path.exists():
        shutil.copy2(path, bak)
    try:
        os.replace(tmp, path)
    except PermissionError as exc:
        raise RuntimeError(
            f"{path} is locked (open in Excel?). Original left intact; new copy kept at {tmp}"
        ) from exc


def check_minimums(counts: dict, minimums: dict) -> None:
    """Raise if any harvest is below its minimum. A missing key counts as 0 (fail closed)."""
    short = {k: (counts.get(k, 0), m) for k, m in minimums.items() if counts.get(k, 0) < m}
    if short:
        raise RuntimeError(f"refusing to save: harvest below minimum {short}")


def cache_sum_formulas(ws, sheet_formula_rows: range | None = None,
                       blank_policy: str = "na") -> dict:
    """Evaluate '=Fr+Gr+Hr' formulas and replace them with values.

    blank_policy:
      "na"   (default, fail closed): any blank operand -> "N/A". A missing
             exchange leg must not silently under-report the total.
      "zero" (explicit opt-in only): blank operand -> 0.
    Numeric text operands are read with num(). Any other text (e.g. "-") -> "N/A".
    Returns {'cached': n, 'na': m, 'skipped': k}.
    NOTE: this replaces the formula. To keep the live formula, write the value
    to a new column instead.
    """
    if blank_policy not in ("na", "zero"):
        raise ValueError("blank_policy must be 'na' or 'zero'")
    out = {"cached": 0, "na": 0, "skipped": 0}
    rows = sheet_formula_rows or range(1, ws.max_row + 1)
    for r in rows:
        for c in ws[r]:
            v = c.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            m = _SUM_FORMULA.match(v)
            if not m:
                out["skipped"] += 1
                continue
            refs = [(m.group(1), int(m.group(2))),
                    (m.group(3), int(m.group(4))),
                    (m.group(5), int(m.group(6)))]
            vals, bad = [], False
            for col, row in refs:
                raw = ws[f"{col}{row}"].value
                if raw is None or (isinstance(raw, str) and raw.strip() == ""):
                    if blank_policy == "zero":
                        vals.append(0.0)
                        continue
                    bad = True
                    break
                x = num(raw)
                if x is None:
                    bad = True
                    break
                vals.append(x)
            if bad:
                c.value = NA
                out["na"] += 1
            else:
                c.value = round(sum(vals), 2)
                out["cached"] += 1
    return out
