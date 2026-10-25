#!/usr/bin/env python3
"""
Terminal/Data_Factory/seed_parquet_from_excel.py
================================================
One-shot bootstrap seeder: Converts all canonical sheets in Data/DataBase_Trading.xlsx
directly into individual production Parquet files in Data/feeds/ and sidecar metadata in Data/metadata/.
"""

from __future__ import annotations

import datetime
import json
import logging
import os
import pathlib
import sys
import time

ROOT_DIR = pathlib.Path(__file__).resolve().parents[2]
DATA_DIR = ROOT_DIR / "Data"
EXCEL_PATH = DATA_DIR / "DataBase_Trading.xlsx"
DATA_FACTORY_DIR = ROOT_DIR / "Terminal" / "Data_Factory"
if str(DATA_FACTORY_DIR) not in sys.path:
    sys.path.insert(0, str(DATA_FACTORY_DIR))

from parquet_feed_writer import write_parquet_feed, ensure_directories

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("ParquetSeeder")

# Ignore orphan legacy sheets
ORPHAN_SHEETS = {"Crypto_ETF_Flows", "Top_Institutions", "ETF_Market_Overview"}


def seed_all_sheets():
    if not EXCEL_PATH.exists():
        log.error(f"Excel file not found at {EXCEL_PATH}")
        return

    import openpyxl
    log.info(f"Loading {EXCEL_PATH} in read_only mode...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, read_only=True)

    ensure_directories()
    seeded_count = 0

    for sheet_name in wb.sheetnames:
        if sheet_name in ORPHAN_SHEETS:
            log.info(f"Skipping orphan sheet: {sheet_name}")
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            log.warning(f"Sheet {sheet_name} is empty.")
            continue

        # Look for banner and header
        # Usually row 0 is Title, row 1 is Banner or Header, row 2 is Header
        # Let's find the first row that looks like table headers
        header_row_idx = None
        for i, r in enumerate(rows[:5]):
            non_empty = [c for c in r if c is not None]
            if len(non_empty) >= 2:
                # Check if it has strings that look like columns
                if all(isinstance(c, str) for c in non_empty[:3]):
                    header_row_idx = i
                    break

        if header_row_idx is None:
            header_row_idx = 0

        raw_headers = rows[header_row_idx]
        headers = []
        for col_idx, h in enumerate(raw_headers):
            if h is not None and str(h).strip():
                h_clean = str(h).strip().replace(" ", "_").replace("(", "").replace(")", "").replace("/", "_")
                headers.append((col_idx, h_clean))
            elif col_idx < 15:
                # auto-name
                headers.append((col_idx, f"col_{col_idx}"))

        data_rows = []
        for r in rows[header_row_idx + 1:]:
            # Skip empty rows
            if not any(c is not None for c in r):
                continue
            row_dict = {}
            for col_idx, col_name in headers:
                val = r[col_idx] if col_idx < len(r) else None
                # Normalize types
                if isinstance(val, (datetime.datetime, datetime.date)):
                    val = str(val)
                elif val is not None and not isinstance(val, (int, float, bool, str)):
                    val = str(val)
                row_dict[col_name] = val
            data_rows.append(row_dict)

        if not data_rows:
            data_rows = [{"_status": "EMPTY_SHEET"}]

        # Write to dedicated parquet file
        meta = write_parquet_feed(
            sheet_name,
            data_rows,
            custom_metadata={"seeded_from": "DataBase_Trading.xlsx", "original_sheet": sheet_name}
        )
        log.info(f"Seeded '{sheet_name}': {meta['row_count']} rows -> Data/feeds/{sheet_name}.parquet")
        seeded_count += 1

    wb.close()
    log.info(f"Successfully seeded {seeded_count} parquet feeds!")


if __name__ == "__main__":
    seed_all_sheets()
